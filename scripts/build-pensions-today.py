#!/usr/bin/env python3
"""Build observed pension distributions; never infer an age × amount joint table.

Inputs are archived official tables in pipeline/source_data/pensions.
Requires openpyxl and xlrd. See pipeline/docs/README_pensions_today.md.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import zipfile
from urllib.parse import quote
from pathlib import Path

import xlrd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "pipeline/source_data/pensions"
OUTPUT = ROOT / "data/pensions-today.v1.json"
EXTRACTED = "2026-09-08"
CSSZ_URL = "https://www.cssz.gov.cz/documents/20143/2946719/Rocenka_2025.zip/a927732b-31f5-7f4c-8487-5ca9223e4dda"
OECD_BASE = "https://www.oecd.org/en/publications/pensions-at-a-glance-2025_e40274c1-en/full-report/"
COUNTRY_NAMES = [
    ("AUS", "Australia", "Austrálie"), ("AUT", "Austria", "Rakousko"), ("BEL", "Belgium", "Belgie"),
    ("CAN", "Canada", "Kanada"), ("CHL", "Chile", "Chile"), ("COL", "Colombia", "Kolumbie"),
    ("CRI", "Costa Rica", "Kostarika"), ("CZE", "Czechia", "Česko"), ("DNK", "Denmark", "Dánsko"),
    ("EST", "Estonia", "Estonsko"), ("FIN", "Finland", "Finsko"), ("FRA", "France", "Francie"),
    ("DEU", "Germany", "Německo"), ("GRC", "Greece", "Řecko"), ("HUN", "Hungary", "Maďarsko"),
    ("ISL", "Iceland", "Island"), ("IRL", "Ireland", "Irsko"), ("ISR", "Israel", "Izrael"),
    ("ITA", "Italy", "Itálie"), ("JPN", "Japan", "Japonsko"), ("KOR", "Korea", "Jižní Korea"),
    ("LVA", "Latvia", "Lotyšsko"), ("LTU", "Lithuania", "Litva"), ("LUX", "Luxembourg", "Lucembursko"),
    ("MEX", "Mexico", "Mexiko"), ("NLD", "Netherlands", "Nizozemsko"), ("NZL", "New Zealand", "Nový Zéland"),
    ("NOR", "Norway", "Norsko"), ("POL", "Poland", "Polsko"), ("PRT", "Portugal", "Portugalsko"),
    ("SVK", "Slovak Republic", "Slovensko"), ("SVN", "Slovenia", "Slovinsko"), ("ESP", "Spain", "Španělsko"),
    ("SWE", "Sweden", "Švédsko"), ("CHE", "Switzerland", "Švýcarsko"), ("TUR", "Türkiye", "Turecko"),
    ("GBR", "United Kingdom", "Spojené království"), ("USA", "United States", "Spojené státy"),
    ("BRA", "Brazil", "Brazílie"), ("UKR", "Ukraine", "Ukrajina"), ("OECD", "OECD average", "Průměr OECD"),
]
YEAR_OVERRIDES = {**dict.fromkeys(["CAN", "CRI", "FIN", "LVA", "NLD", "SWE", "GBR", "USA"], 2023),
                  "DEU": 2021, "JPN": 2021, "AUS": 2020, "ISL": 2017}


def oecd(data):
    names = {en: code for code, en, _ in COUNTRY_NAMES}
    names["OECD"] = "OECD"
    for code, en, cs in COUNTRY_NAMES:
        data["countries"].setdefault(code, {"name": {"en": en, "cs": cs}})
    for metric, file, sheet, start, end, indices, link, page in [
        ("income", "oecd-income.xlsx", "t7-1", 7, 26, {"all":1, "male":3, "female":4, "age_66_75":5, "age_76_plus":6}, "crov86", "incomes-of-older-people_24a65c56.html"),
        ("poverty", "oecd-poverty.xlsx", "t7-2", 8, 27, {"all":1, "age_66_75":2, "age_76_plus":3, "male":4, "female":5, "total_population":6}, "2sqwtk", "old-age-income-poverty_3cd946d7.html"),
    ]:
        workbook = load_workbook(INPUT / file, read_only=True, data_only=True)
        rows = list(workbook[sheet].values)
        data["sources"][metric] = source(file, OECD_BASE + page, f"OECD Pensions at a Glance 2025 — Table {sheet[1:].replace('-', '.')}", download_url=f"https://stat.link/{link}", year=2025)
        for row in rows[start-1:end]:
            for offset in [0, 7]:
                label = row[offset]
                if label is None:
                    continue
                label = " ".join(label.split())
                code = names[label]
                if row[offset+1] is None:
                    data["countries"][code][metric] = None
                    continue
                record = {key: round(row[offset+i], 6) for key, i in indices.items()}
                record.update(year=YEAR_OVERRIDES.get(code, 2022) if code != "OECD" else None, source_id=metric)
                data["countries"][code][metric] = record
        workbook.close()
    assert data["countries"]["COL"]["income"] is None
    assert round(data["countries"]["CZE"]["income"]["age_66_75"], 1) == 80.4
    assert round(data["countries"]["CZE"]["poverty"]["age_76_plus"], 1) == 8.9


def source(file, url, table, **extra):
    return {"url": url, "table": table, "extracted_at": EXTRACTED,
            "vintage": "outturn", "file": file,
            "sha256": hashlib.sha256((INPUT / file).read_bytes()).hexdigest(), **extra}


def bounds(label):
    text = str(label).replace(" ", "").replace("\xa0", "")
    nums = [int(n) for n in re.findall(r"\d+", text)]
    if not nums:
        return None
    return nums[0], None if "+" in text else nums[-1]


def czech():
    archive = zipfile.ZipFile(INPUT / "cssz-2025.zip")
    book = xlrd.open_workbook(file_contents=archive.read(next(n for n in archive.namelist() if n.startswith("07.03"))))
    sexes = {}
    for sex, sheet in [("total", "S-celkem"), ("male", "S-muži"), ("female", "S-ženy")]:
        rows = [book.sheet_by_name(sheet).row_values(i) for i in range(book.sheet_by_name(sheet).nrows)]
        total = next(r for r in rows if r[0] == "ÚHRN")
        mean = next(r for r in rows if r[0] == "PRŮM. VÝŠE")
        start = next(i for i, r in enumerate(rows) if r[0] == "MĚSÍČNÍ VÝŠE") + 1
        values = []
        for r in rows[start:]:
            if r[0] == "NEUDÁNO":
                assert all(v == 0 for v in r[1:]), "Unknown amounts must be preserved"
                break
            lo, hi = bounds(r[0])
            values.append((lo, hi, r))
        types = {}
        for kind, col in [("all", 7), ("regular", 1), ("early", 4)]:
            bands = [{"lower": lo, "upper": hi, "count": int(r[col])} for lo, hi, r in values]
            assert sum(b["count"] for b in bands) == total[col]
            count = int(total[col])
            cumulative = 0
            median_band = None
            for b in bands:
                cumulative += b["count"]
                if median_band is None and cumulative >= count / 2:
                    median_band = {"lower": b["lower"], "upper": b["upper"]}
            types[kind] = {"count": count, "mean": mean[col], "median_band": median_band, "bands": bands}
        sexes[sex] = types
    for kind in sexes["total"]:
        assert sexes["total"][kind]["count"] == sum(sexes[s][kind]["count"] for s in ["male", "female"])
        for i, b in enumerate(sexes["total"][kind]["bands"]):
            assert b["count"] == sum(sexes[s][kind]["bands"][i]["count"] for s in ["male", "female"])
    return {"date": "2025-12-31", "currency": "CZK", "frequency": "month", "source_id": "cssz",
            "distribution": sexes, "payment_by_age": None, "payment_by_retirement_year": None}


def national_extensions(data):
    file = "ssa-2026-table-5a11.csv"
    rows = list(csv.DictReader((INPUT / file).open()))
    sexes = {}
    for sex in ["total", "male", "female"]:
        count = int(rows[0][f"{sex}_count"])
        mean = float(rows[0][f"{sex}_mean"])
        ages = [{"age": r["age"], "count": int(r[f"{sex}_count"]), "mean": float(r[f"{sex}_mean"])} for r in rows[1:]]
        assert sum(r["count"] for r in ages) == count
        assert abs(sum(r["count"] * r["mean"] for r in ages) / count - mean) < 0.011
        sexes[sex] = {"count": count, "mean": mean, "ages": ages}
    for r in rows:
        assert int(r["total_count"]) == int(r["male_count"]) + int(r["female_count"])
    data["sources"]["ssa"] = source(file, "https://www.ssa.gov/policy/docs/statcomps/supplement/2026/5a.html#table5.a1.1", "SSA Annual Statistical Supplement 2026 — Table 5.A1.1", year=2026, extraction_method="Transcribed published age-band rows; counts and weighted means reconciled to published totals")
    data["countries"]["USA"]["national"] = {"date": "2025-12-31", "currency": "USD", "frequency": "month", "source_id": "ssa", "payment_by_age": sexes, "payment_by_retirement_year": None}

    file = "dwp-2026-february-state-pension.csv"
    rows = list(csv.DictReader((INPUT / file).open()))
    data["sources"]["dwp"] = source(file, "https://www.gov.uk/government/statistics/dwp-benefit-statistics-february-2026/dwp-benefit-statistics-february-2026#pensions", "DWP benefit statistics, February 2026 — Figure 4 and accompanying means", year=2026, extraction_method="Transcribed published scheme/sex means")
    data["countries"]["GBR"]["national"] = {"date": "2025-08", "currency": "GBP", "frequency": "week", "source_id": "dwp", "payment_by_scheme": {sex: {r["scheme"]: float(r[f"{sex}_mean"]) for r in rows} for sex in ["total", "male", "female"]}, "payment_by_retirement_year": None}

    file = "drees-2025-05.xlsx"
    w = load_workbook(INPUT / file, data_only=True, read_only=True)
    latest = next(r for r in w["F05_Tableau 1"].values if r[1] == 2023)
    assert latest[5] == 1666 and latest[6] == 1541
    p = {"date": "2023-12-31", "currency": "EUR", "frequency": "month", "source_id": "drees_level",
         "means": {"total": {"direct": latest[5], "combined": latest[9]}, "male": {"direct": latest[4], "combined": latest[8]}, "female": {"direct": latest[3], "combined": latest[7]}},
         "payment_by_retirement_year": None}
    data["sources"]["drees_level"] = source(file, "https://www.drees.solidarites-sante.gouv.fr/sites/default/files/2025-07/" + quote("Fiche 05 - Le niveau des pensions.xlsx"), "DREES 2025 — Fiche 05, Tableau 1 / Graphique 4", year=2025)
    distribution = list(w["F05_Graphique 4"].values)
    p["distribution_pct"] = {}
    for sex, column in [("total",4),("male",3),("female",2)]:
        bands = [{"upper": 100 if i == 4 else None if i == 49 else distribution[i][1], "share": distribution[i][column]} for i in range(4,50)]
        assert abs(sum(b["share"] for b in bands) - 100) < 1
        p["distribution_pct"][sex] = bands
    p["distribution_date"] = "2020-12-31"
    w.close()
    file = "drees-2025-06.xlsx"
    w = load_workbook(INPUT / file, data_only=True, read_only=True)
    # Keep only cohorts observed in the common 2020 snapshot. The later 1954–56
    # rows use ANCETRE 2021–23 and revaluation adjustments; never blend them here.
    p["payment_by_birth_year"] = {"total": [{"birth_year": r[1], "mean": r[6], "full_career_mean": r[8]} for r in w["F06 - Graphique 1"].values if isinstance(r[1], int) and 1930 <= r[1] <= 1953]}
    for sex, mean_col, full_col in [("female",5,7),("male",11,13)]:
        p["payment_by_birth_year"][sex] = [{"birth_year":r[1],"mean":r[mean_col],"full_career_mean":r[full_col]} for r in w["F06 - Graphique 1 compl"].values if isinstance(r[1],int) and 1930 <= r[1] <= 1953]
    assert all(len(rows) == 24 for rows in p["payment_by_birth_year"].values())
    p["birth_year_date"] = "2020-12-31"
    w.close()
    data["sources"]["drees_cohort"] = source(file, "https://www.drees.solidarites-sante.gouv.fr/sites/default/files/2025-07/" + quote("Fiche 06 - Les écarts de pensions de droit direct entre générations.xlsx"), "DREES 2025 — Fiche 06, Graphique 1 / complémentaire, resident cohorts 1930–1953", year=2025)
    data["countries"]["FRA"]["national"] = p


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--stage", choices=["czech", "all"], default="all")
    args = parser.parse_args()
    data = {"schema_version": 1, "extracted_at": EXTRACTED,
            "sources": {"cssz": source("cssz-2025.zip", CSSZ_URL, "07.03 — S-celkem, S-muži, S-ženy", year=2025)},
            "countries": {"CZE": {"name": {"en": "Czechia", "cs": "Česko"}, "national": czech()}}}
    if args.stage == "all":
        oecd(data)
        national_extensions(data)
    import runpy
    data = runpy.run_path(str(ROOT / "scripts/build-pension-awards.py"))["build"](data)
    OUTPUT.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n")
    print(f"Built {OUTPUT.name}: Czech pension distributions reconcile to ČSSZ totals; "
          f"{sum(bool(c.get('income')) for code,c in data['countries'].items() if code != 'OECD')} international age-income profiles")


if __name__ == "__main__":
    main()
