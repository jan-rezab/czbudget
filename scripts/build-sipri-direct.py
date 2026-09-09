"""Extract Czech source observations from the frozen SIPRI April 2026 workbook.

Keep full precision, source cells, notes and font colours (SIPRI estimate flags).
The Czech GDP-share series replaces the older WDI vintage as a whole.
"""
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
PATH = ROOT.parent / 'data/sources/defense/SIPRI-Milex-data-1949-2025_v1.2.xlsx'
URL = 'https://www.sipri.org/sites/default/files/SIPRI-Milex-data-1949-2025_v1.2.xlsx'

def main():
    book = load_workbook(PATH, read_only=True, data_only=True)
    series = []
    for sheet in book:
        rows = list(sheet.iter_rows())
        header = next((r for r in rows if r[0].value == 'Country'), None)
        country = next((r for r in rows if r[0].value == 'Czechia'), None)
        if not header or not country:
            continue
        observations = []
        for head, cell in zip(header, country):
            if not isinstance(head.value, int) or not 1949 <= head.value <= 2025:
                continue
            color = cell.font.color
            observations.append({'year': head.value, 'value': cell.value if isinstance(cell.value, (int, float)) else None,
                'source_value': cell.value, 'cell': cell.coordinate, 'number_format': cell.number_format,
                'font_color': {'type': color.type, 'value': str(color.value)} if color else None})
        series.append({'sheet': sheet.title, 'title': rows[0][0].value, 'country': 'CZE',
            'notes': country[1].value, 'observations': observations})
    notes = [list(r) for r in book['Footnotes'].values if r[0] == 71]
    source = {'provider': 'SIPRI', 'title': 'SIPRI Military Expenditure Database, April 2026 revision',
        'url': URL, 'edition': '1949–2025 v1.2; revised 27 April 2026', 'retrieved_at': '2026-09-09',
        'sha256': hashlib.sha256(PATH.read_bytes()).hexdigest(), 'attribution': '© SIPRI 2026. Direct-source terms apply; WDI licence is not inherited.'}
    output = {'schema_version': '1.0.0', 'generated_at': datetime.now(timezone.utc).isoformat(),
        'source': source, 'notes': notes, 'flag_legend': 'Blue font: SIPRI estimate. Red font: highly uncertain. Source font colours retained.', 'series': series}
    (ROOT / 'data/czech-sipri-military-expenditure.v1.json').write_text(json.dumps(output, ensure_ascii=False, indent=2) + '\n')
    path = ROOT / 'data/defense-deep-dive.v1.json'
    payload = json.loads(path.read_text())
    czech = next(c for c in payload['countries'] if c['code'] == 'CZE')
    points = next(s for s in series if s['sheet'] == 'Share of GDP')['observations']
    values = [[p['year'], p['value'] * 100] for p in points if p['value'] is not None]
    assert values[-1][0] == 2025 and 1 < values[-1][1] < 3
    czech['comparison'] = {'indicator': 'SIPRI_SHARE_GDP', 'unit': 'pct_gdp', 'series': values,
        'latest': {'year': values[-1][0], 'value': values[-1][1]}, 'source': source,
        'method': 'Percentage-formatted source fractions multiplied by 100. Full revised history replaces WDI vintage.'}
    payload['generated_at'] = output['generated_at']
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + '\n')
    print(f'Czech SIPRI: {len(series)} series; latest GDP share {values[-1]}')

if __name__ == '__main__':
    main()
