#!/usr/bin/env python3
"""MF actual-year consolidated accounting, separate from unconsolidated entity sums."""
import json,hashlib,importlib.util
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
CACHE=ROOT.parent/'data/source_cache/mf_consolidated'
BASE='https://mf.gov.cz/assets/attachments/'
FILES={'history-2024.xlsx':'2024-12-31_Priloha-c-2-Vyvoj-jednotlivych-polozek-ucetnich-vykazu-za-Ceskou-republiku-2016-2024.xlsx','perimeter-2024.xlsx':'2024-12-31_Priloha-c-3-Konsolidacni-celek-Ceska-republika.xlsx'}
def extract_history(w):
    out=[]
    for s in w:
        rows=list(s.values);year=None;cols=[]
        assets=s.title.endswith('AKTIVA')
        for i,v in enumerate(rows[0]):
            if isinstance(v,int) and 2016<=v<=2024:year=v
            if i>=3 and year and (assets or isinstance(v,int)):
                cols.append((i,year,rows[1][i] if assets else 'reported'))
        for n,r in enumerate(rows[2:] if assets else rows[1:],3 if assets else 2):
            if not any(r[:3]):continue
            code=r[1] or r[0];label=r[2] or code
            for i,y,measure in cols:
                if i<len(r) and isinstance(r[i],(int,float)):
                    out.append({'statement':s.title,'code':str(code),'label_cs':str(label),'year':y,'measure':measure,'value_m_czk':r[i],'source_row':n,'source_column':i+1})
    return out

def main():
    hist=extract_history(load_workbook(CACHE/'history-2024.xlsx',read_only=True,data_only=True))
    assets={r['year']:r['value_m_czk'] for r in hist if r['code']=='AKTIVA' and r['measure']=='Netto'}
    liabilities={r['year']:r['value_m_czk'] for r in hist if r['code']=='PASIVA'}
    if set(assets)!=set(range(2016,2025)) or any(abs(assets[y]-liabilities[y])>0.01 for y in assets):raise ValueError('Consolidated balance sheet does not reconcile')
    w=load_workbook(CACHE/'perimeter-2024.xlsx',read_only=True,data_only=True);members=[]
    for s in w:
        for n,r in enumerate(s.values,1):
            if (isinstance(r[0],(int,float)) and r[0]>0 or isinstance(r[0],str) and r[0].strip().isdigit()) and len(r)>1 and isinstance(r[1],str):
                members.append({'ico':str(int(r[0])).zfill(8),'name':r[1],'source_sheet':s.title,'source_row':n,'source_cells':[str(v) if v is not None else None for v in r],'period':2024})
    payload={'schema_version':'1.0.0','period':{'from':2016,'to':2024},'unit':'CZK_million','basis':'accrual_consolidated','scope':'MF consolidation perimeter; not S13 or state budget cash','aggregation':'Overlapping subtotal and leaf rows. Do not sum all rows. Reported profit includes consolidation adjustments.','history':hist,'perimeter':members,'perimeter_source_sheets':[{'name':s.title,'rows':[[str(v) if v is not None else None for v in row] for row in s.values if any(v is not None for v in row)]} for s in w],'historical_perimeter_detail':'data/czech-mf-perimeter-history.v1.json','perimeter_note':'Source rows include annual changes; raw change annotations retained, not asserted as current active count','sources':[{'url':BASE+url,'sha256':hashlib.sha256(((CACHE/name)).read_bytes()).hexdigest()} for name,url in FILES.items()]}
    (ROOT/'data/czech-consolidated-accounts.v1.json').write_text(json.dumps(payload,ensure_ascii=False,separators=(',',':'))+'\n')
    print(json.dumps({'observations':len(hist),'perimeter_rows':len(members),'asset_balance_checks':len(assets)}))
if __name__=='__main__':main()
