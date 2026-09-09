#!/usr/bin/env python3
"""Preserve the official operation/procurement row grain; never sum repeated project amounts."""
import gzip,json,hashlib
from pathlib import Path
from datetime import date,datetime
import openpyxl
ROOT=Path(__file__).resolve().parents[2];OUT=ROOT/'website/data';CACHE=ROOT/'data/source_cache/dotaceeu'
URL='https://www.dotaceeu.cz/getmedia/e06f478c-d716-4dac-bfd7-48a8b6cc0d6b/2026_08_Seznam-operaci_List-of-Operations_21.xlsx.aspx?ext=.xlsx'
def scalar(v):return v.isoformat() if isinstance(v,(date,datetime)) else v
def main():
 p=CACHE/'operations-202608.xlsx';s=openpyxl.load_workbook(p,read_only=True,data_only=True).active
 rows=s.iter_rows(values_only=True);title=next(rows);generated=next(rows);next(rows);headers=list(next(rows));next(rows)
 assert headers[7]=='Registrační číslo projektu' and headers[26]=='Celkové náklady na operaci (CZK)' and len(headers)==48
 dest=OUT/'dotaceeu/operation-rows.ndjson.gz';dest.parent.mkdir(exist_ok=True);tmp=dest.with_suffix('.tmp')
 ids=set();count=0;proc=0
 with gzip.open(tmp,'wt',encoding='utf-8') as f:
  for n,row in enumerate(rows,6):
   if not row[7]:continue
   if not str(row[7]).startswith('CZ.'):raise ValueError(f'Unexpected operation ID at {n}')
   count+=1;ids.add(row[7]);proc+=bool(row[36])
   obj={'project_id':row[7],'source_row':n,'beneficiary_ico':str(row[14]) if row[14] is not None else None,'status_code':row[11],'source_values':[scalar(v) for v in row]}
   f.write(json.dumps(obj,ensure_ascii=False,separators=(',',':'))+'\n')
 tmp.replace(dest)
 meta={'schema_version':1,'country':'CZE','source_url':URL,'source_sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'source_generated':generated[0],'as_of':'2026-08-01','programming_period':'2021-2027','source_columns':headers,'row_count':count,'unique_project_count':len(ids),'rows_with_procurement_title':proc,'detail_path':'data/dotaceeu/operation-rows.ndjson.gz','grain':'Source operation/procurement row; project IDs can repeat.','financial_stages':{'26':'total_operation_cost','27-30':'legal_act_commitment','31-34':'charged_in_payment_application_not_cash_paid','45':'procurement_estimate_ex_vat','46':'contract_price_ex_vat','47':'actually_paid_project_procurement_ex_vat'},'scope_notes':['Only operations with a legal act; rural development programme excluded.','Do not sum operation financial columns across repeated procurement rows.','Do not combine with ReD or MONITOR without explicit cross-source deduplication.','All48 original columns retained; null does not mean zero.']}
 annex=CACHE/'instruments-202608.xlsx'
 if annex.exists():
  book=openpyxl.load_workbook(annex,read_only=True,data_only=True)
  meta['financial_instruments_annex']={'source_url':'https://www.dotaceeu.cz/getmedia/64ac77fd-33eb-4362-9313-ac07426dc49e/2026_08_Seznam-operaci-FN_2021_2027.xlsx.aspx?ext=.xlsx','source_sha256':hashlib.sha256(annex.read_bytes()).hexdigest(),'as_of':'2026-08-04','scope':'Separate financial-instrument operations annex; no combined totals with main list.','sheets':[{'sheet':sheet.title,'rows':[[scalar(v) for v in row] for row in sheet.iter_rows(values_only=True)]} for sheet in book]}
 (OUT/'czech-dotaceeu-operations.v1.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2)+'\n');print(count,len(ids),proc)
if __name__=='__main__':main()
