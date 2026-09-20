#!/usr/bin/env python3
"""Build the Latin America million-city anchor registry from official UN WUP F21.

This is a crosswalk, not a spatial aggregation. A published match means PSD has
line items for the named legal-government anchor; it never means that every
government inside the UN Degree-of-Urbanization city is consolidated.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
CONFIG = ROOT / "pipeline/config/latam_million_cities_sources.json"
PROFILES = ROOT / "data/municipal-expansion"
COVERAGE = ROOT / "data/municipal-itemized-coverage.v1.json"
LAC = {"ARG", "BOL", "BRA", "CHL", "COL", "CRI", "CUB", "DOM", "ECU", "SLV", "GTM", "HTI", "HND", "MEX", "NIC", "PAN", "PRY", "PER", "URY", "VEN"}


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            value.update(chunk)
    return value.hexdigest()


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--un-workbook", required=True, type=Path)
    parser.add_argument("--output", type=Path, default=ROOT / "data/major-cities/latam-million-plus.v1.json")
    args = parser.parse_args()

    config = json.loads(CONFIG.read_text(encoding="utf-8"))
    actual_hash = digest(args.un_workbook)
    if actual_hash != config["un_source"]["sha256"]:
        raise SystemExit(f"UN workbook hash mismatch: {actual_hash}")

    coverage_rows = json.loads(COVERAGE.read_text(encoding="utf-8"))["countries"]
    coverage = {row["code"]: row for row in coverage_rows}
    workbook = load_workbook(args.un_workbook, read_only=True, data_only=True)
    sheet = workbook["Data"]
    header = next(sheet.iter_rows(values_only=True))
    year_column = header.index("2025")
    records = []

    for row in sheet.iter_rows(min_row=2, values_only=True):
        country, population = row[4], row[year_column]
        if country not in LAC or population is None or population < 1000:
            continue
        city_code = str(row[6])
        mapping = config["anchor_entities"].get(city_code)
        blocker = config["city_blockers"].get(city_code) or config["country_blockers"].get(country)
        item = {
            "un_city_code": int(row[6]),
            "un_city_name": row[7],
            "country_code": country,
            "population_2025_thousands": population,
            "longitude": row[9],
            "latitude": row[10],
            "coverage_status": "blocked",
            "coverage_scope": "UN Degree-of-Urbanization city",
            "fiscal_scope": None,
            "profile_url": None,
            "source_url": None,
            "period": None,
            "stages": [],
            "blocker": blocker or "No reviewed anchor mapping",
        }
        if mapping:
            layer = coverage[country]
            if layer["status"] not in {"full", "partial"}:
                raise SystemExit(f"{country} mapping points to non-itemized layer: {layer['status']}")
            profile_path = PROFILES / country.lower() / f"{mapping['entity_code']}.json"
            if not profile_path.exists():
                raise SystemExit(f"Missing mapped profile: {profile_path}")
            profile = json.loads(profile_path.read_text(encoding="utf-8"))
            if len(profile.get("detail", [])) == 0:
                raise SystemExit(f"Mapped profile has no detail rows: {profile_path}")
            item.update({
                "coverage_status": "anchor_only",
                "fiscal_scope": mapping["fiscal_boundary"],
                "profile_entity_code": mapping["entity_code"],
                "profile_entity_name": profile["name"],
                "profile_url": profile["url"],
                "source_url": layer["source_url"],
                "period": layer["period"],
                "stages": layer["stages"],
                "published_detail_rows": len(profile["detail"]),
                "blocker": "Full agglomeration coverage is not measured; this profile covers only the named fiscal entity",
            })
        override = config.get("warehouse_city_overrides", {}).get(city_code)
        if override:
            item.pop("profile_entity_code", None)
            item.pop("profile_entity_name", None)
            item.pop("published_detail_rows", None)
            item.update({
                "coverage_status": "anchor_only",
                "fiscal_scope": override["fiscal_boundary"],
                "profile_url": None,
                "source_url": override["source_url"],
                "period": override["period"],
                "stages": override["stages"],
                "blocker": "Full agglomeration coverage is not measured; warehouse facts cover only the named fiscal entity",
                "warehouse_entity_id": override["entity_id"],
                "warehouse_entity_name": override["entity_name"],
                "warehouse_status": override["warehouse_status"],
                "warehouse_rows": override["warehouse_rows"],
                "source_id": override["source_id"],
                "completed_run": override["completed_run"],
                "ingestion_run_id": override["ingestion_run_id"],
            })
        records.append(item)

    records.sort(key=lambda item: (-item["population_2025_thousands"], item["country_code"], item["un_city_name"]))
    counts = {
        "un_cities": len(records),
        "anchor_only": sum(item["coverage_status"] == "anchor_only" for item in records),
        "blocked": sum(item["coverage_status"] == "blocked" for item in records),
        "full_agglomeration": 0,
    }
    if counts != {"un_cities": 60, "anchor_only": 46, "blocked": 14, "full_agglomeration": 0}:
        raise SystemExit(f"Unexpected coverage counts: {counts}")
    payload = {
        "schema_version": "1.0.0",
        "generated_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "source": {**config["un_source"], "retrieved_sha256": actual_hash},
        "methodology": "A city is anchor_only only when its reviewed legal-government profile exists and its country passes PSD's five-code itemization floor, or when a reviewed city-specific warehouse override records a completed itemized load. No city is claimed as full agglomeration coverage.",
        "counts": counts,
        "cities": records,
        "next_official_sources": config["next_official_sources"],
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(counts))


if __name__ == "__main__":
    main()
