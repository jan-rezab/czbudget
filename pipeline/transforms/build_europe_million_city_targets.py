#!/usr/bin/env python3
"""Build the audited European million-city target bundle from UN WUP F21.

This is intentionally a crosswalk, not an urban-area aggregation. Each matched
record points at one legal fiscal entity and carries that limitation in output.
"""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


EUROPE_ISO3 = {
    "ALB", "AND", "AUT", "BLR", "BEL", "BIH", "BGR", "HRV", "CYP", "CZE",
    "DNK", "EST", "FIN", "FRA", "DEU", "GRC", "HUN", "ISL", "IRL", "ITA",
    "LVA", "LIE", "LTU", "LUX", "MLT", "MDA", "MCO", "MNE", "NLD", "MKD",
    "NOR", "POL", "PRT", "ROU", "RUS", "SMR", "SRB", "SVK", "SVN", "ESP",
    "SWE", "CHE", "UKR", "GBR", "VAT", "KOS",
}


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def load_directory(root: Path, iso3: str) -> dict[str, dict]:
    path = root / iso3.lower() / "municipalities.v1.json"
    if not path.exists():
        return {}
    payload = json.loads(path.read_text(encoding="utf-8"))
    return {row["id"]: row for row in payload.get("entities", [])}


def public_artifact(root: Path, entity: dict) -> tuple[str | None, int | None]:
    code = entity["code"]
    country = entity["country"].lower()
    candidates = [
        root / "municipal-expansion" / country / f"{code}.json",
        root / "municipal-benchmarks" / country / f"{code}.json",
    ]
    for path in candidates:
        if not path.exists():
            continue
        payload = json.loads(path.read_text(encoding="utf-8"))
        lines = payload.get("detail") or payload.get("breakdown") or []
        return path.relative_to(root).as_posix(), len(lines)
    return None, None


def build(workbook: Path, config_path: Path, country_root: Path, data_root: Path) -> dict:
    config = json.loads(config_path.read_text(encoding="utf-8"))
    targets = config["targets"]
    sheet = load_workbook(workbook, read_only=True, data_only=True)["Data"]
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))
    year_index = header.index("2025")
    records = []
    directories: dict[str, dict[str, dict]] = {}

    for row in rows:
        population = row[year_index]
        iso3 = row[4]
        if iso3 not in EUROPE_ISO3 or population is None or population < 1000:
            continue
        city_code = str(row[6])
        target = targets.get(city_code)
        record = {
            "un_city_code": int(row[6]),
            "un_city_name": row[7],
            "country": row[1],
            "iso3": iso3,
            "population_2025": int(round(float(population) * 1000)),
            "un_centroid": {"longitude": row[9], "latitude": row[10]},
            "fiscal_scope": "anchor_legal_government_only" if target else None,
            "status": target["status"] if target else "no_adapter",
        }
        if target:
            directories.setdefault(iso3, load_directory(country_root, iso3))
            entity = directories[iso3].get(target["anchor_entity_id"])
            if entity is None:
                raise ValueError(f"{city_code}: missing anchor {target['anchor_entity_id']}")
            if entity["name"] != target["anchor_name"]:
                raise ValueError(f"{city_code}: anchor name drift: {entity['name']!r}")
            artifact, line_count = public_artifact(data_root, entity)
            record.update({
                "anchor_entity_id": entity["id"],
                "anchor_entity_code": entity["code"],
                "anchor_entity_name": entity["name"],
                "anchor_years": entity.get("years", []),
                "profile_url": entity.get("url"),
                "public_artifact": artifact,
                "public_line_count": line_count,
                "source_ids": target["source_ids"],
                "scope_note": target.get("scope_note", "Anchor municipality only; the UN built-up area may include other governments."),
            })
            evidence = config.get("warehouse_verification", {}).get("entities", {}).get(entity["id"])
            if evidence:
                record["warehouse_verification"] = evidence
        records.append(record)

    if len(records) != 45:
        raise ValueError(f"Expected 45 European million-plus cities, found {len(records)}")
    if set(targets) - {str(row["un_city_code"]) for row in records}:
        raise ValueError("Configuration contains a target outside the selected UN universe")

    counts = Counter(row["status"] for row in records)
    return {
        "schema_version": "1.0.0",
        "dataset_id": config["dataset_id"],
        "generated_at": date.today().isoformat(),
        "source": {**config["un_source"], "sha256": sha256(workbook)},
        "methodology": config["methodology"],
        "warehouse_verification": {key: value for key, value in config.get("warehouse_verification", {}).items() if key != "entities"},
        "summary": {
            "european_un_cities_over_1m": len(records),
            "anchor_targets": sum(counts[s] for s in config["methodology"]["loadable_statuses"]),
            "status_counts": dict(sorted(counts.items())),
        },
        "cities": sorted(records, key=lambda row: (-row["population_2025"], row["un_city_name"])),
    }


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", type=Path, required=True)
    parser.add_argument("--config", type=Path, default=Path("pipeline/config/europe_million_city_targets.json"))
    parser.add_argument("--country-root", type=Path, default=Path("data/countries"))
    parser.add_argument("--data-root", type=Path, default=Path("data"))
    parser.add_argument("--output", type=Path, default=Path("data/europe-million-city-budget-targets.v1.json"))
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    payload = build(args.workbook, args.config, args.country_root, args.data_root)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload["summary"], indent=2))


if __name__ == "__main__":
    main()
