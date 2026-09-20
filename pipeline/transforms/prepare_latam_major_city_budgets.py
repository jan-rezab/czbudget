#!/usr/bin/env python3
"""Normalize reviewed Latin-American city budget files into warehouse JSONL.

Bulk downloads and full transformations belong on the cloud worker. Locally,
use only small synthetic fixtures or an explicitly supplied temporary file.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import zipfile
from openpyxl import load_workbook
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2] if len(Path(__file__).resolve().parents) > 2 else Path.cwd()
CONFIG = ROOT / "pipeline/config/latam_major_city_budget_sources.json"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def amount(value: str | None) -> str | None:
    text = (value or "").strip().replace(".", "").replace(",", ".")
    if not text:
        return None
    try:
        number = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError(f"Invalid source amount: {value!r}") from exc
    if not number:
        return None
    return format(number.quantize(Decimal("0.000001")), "f")


def rows_from_zip(path: Path, member_prefix: str):
    with zipfile.ZipFile(path) as archive:
        members = [name for name in archive.namelist() if Path(name).name.lower().startswith(member_prefix)]
        if len(members) != 1:
            raise ValueError(f"Expected one {member_prefix} CSV, found {members}")
        with archive.open(members[0]) as raw:
            yield members[0], csv.DictReader(io.TextIOWrapper(raw, encoding="utf-8-sig", newline=""))


def write_jsonl(path: Path, rows: list[dict]) -> None:
    path.write_text("".join(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n" for row in rows), encoding="utf-8")


def spanish_amount(value: str | None) -> str | None:
    text = (value or "").strip().replace(".", "").replace(",", ".")
    if not text:
        return None
    try:
        number = Decimal(text)
    except InvalidOperation:
        return None
    return None if not number else format(number.quantize(Decimal("0.000001")), "f")


def buenos_aires(source: dict, archive: Path, output: Path, loaded_at: str) -> dict:
    actual_hash = sha256(archive)
    if actual_hash != source["verified_sha256"]:
        raise ValueError(f"Source hash mismatch: {actual_hash}")
    reader = csv.DictReader(io.StringIO(archive.read_text(encoding="latin-1")), delimiter=";")
    required = {"Car", "Jur", "UE", "Prog", "Fin", "Fun", "Inciso", "Ppal", "Parc", "Sparc", "Eco", "Fte", "Geo", *source["stage_mapping"]}
    if not reader.fieldnames or not required.issubset(reader.fieldnames):
        raise ValueError(f"Unexpected Buenos Aires columns: {reader.fieldnames}")
    facts=[]; rows_read=0
    dimension_codes=("Car","Jur","Sjur","Ent","Ogese","UE","Prog","Sprog","Proy","Act","Obra","Fin","Fun","Inciso","Ppal","Parc","Sparc","Eco","Fte","Geo")
    for row_number,row in enumerate(reader,2):
        if not any((value or "").strip() for value in row.values() if isinstance(value,str)): continue
        if None in row: raise ValueError(f"Malformed semicolon CSV row {row_number}")
        rows_read += 1
        composite="|".join((row.get(key) or "0").strip() for key in dimension_codes)
        economic=".".join((row.get(key) or "0").strip() for key in ("Inciso","Ppal","Parc","Sparc","Eco"))
        label=" / ".join(filter(None, ((row.get(key) or "").strip() for key in ("Desc_Inc","Desc_Ppal","Desc_Parc","Desc_Sparc","Desc_Eco"))))
        for column,stage in source["stage_mapping"].items():
            value=spanish_amount(row.get(column))
            if value is None: continue
            facts.append({"public_entity_id":"AR:CABA","fiscal_year":2025,"fiscal_period":"FY","reporting_scope":"standalone_autonomous_city","budget_stage":stage,"budget_side":"expenditure","source_budget_item_type_code":label or economic,"functional_paragraph_code":"|".join(((row.get("Fin") or "0").strip(),(row.get("Fun") or "0").strip())),"economic_item_code":economic+"|"+composite,"amount_local":value,"currency_code":"ARS","amount_eur":None,"fx_date":None,"is_consolidation_item":False,"is_financing":False,"is_summary_row":False,"source_row_number":row_number,"source_sheet":"Presupuesto Ejecutado 2025 Q4","source_id":source["id"],"ingestion_run_id":source["id"],"quality_flags":["official_city_open_data","maximum_detail_execution","anchor_only_not_un_agglomeration"],"loaded_at":loaded_at})
    output.mkdir(parents=True,exist_ok=True)
    entities=[{"public_entity_id":"AR:CABA","entity_name":"Gobierno de la Ciudad Autónoma de Buenos Aires","entity_type":"autonomous_city","country_code_alpha2":"AR","country_code_alpha3":"ARG","national_entity_code":"CABA","national_entity_code_type":"AR_AUTONOMOUS_CITY_CODE","is_eu_capital":False,"is_extra_city":True,"default_currency_code":"ARS","eurostat_city_code":None,"eurostat_geography_name":None,"administrative_region_code":"CABA","administrative_region_name":"Ciudad Autónoma de Buenos Aires","administrative_district_code":None,"administrative_district_name":None,"national_geography_code":"CABA","national_geography_code_type":"AR_AUTONOMOUS_CITY_CODE","valid_from":None,"valid_to":None,"loaded_at":loaded_at}]
    sources=[{"source_id":source["id"],"public_entity_id":"AR:CABA","source_type":"official_city_open_data","source_name":"Presupuesto Ejecutado 2025 Cuarto Trimestre","source_url":source["resource_url"],"dataset_code":"presupuesto-ejecutado-2025-q4","archive_file":archive.name,"archive_sha256":actual_hash,"retrieved_at":source["retrieved_at"],"notes":source["fiscal_boundary"]+"; official landing="+source["landing_url"],"loaded_at":loaded_at}]
    runs=[{"ingestion_run_id":source["id"],"source_id":source["id"],"started_at":loaded_at,"completed_at":loaded_at,"status":"completed","source_vintage":"2025-Q4","source_sha256":actual_hash,"rows_read":rows_read,"rows_loaded":len(facts),"warning_count":1,"error_message":"Anchor jurisdiction only; no metropolitan consolidation."}]
    for name,records in (("public_entities",entities),("public_entity_sources",sources),("municipal_budget_line_facts",facts),("ingestion_runs",runs)): write_jsonl(output/f"{name}.jsonl",records)
    manifest={"source_id":source["id"],"source_sha256":actual_hash,"rows_read":rows_read,"facts":len(facts),"entity_count":1}
    (output/"manifest.json").write_text(json.dumps(manifest,indent=2)+"\n")
    (output/"international_municipal_manifest.json").write_text(json.dumps({"schema_version":"1.0.0","status":"validated","countries":["ARG"],"sources":[source["id"]],"counts":{"public_entities":1,"public_entity_sources":1,"municipal_budget_line_facts":len(facts),"ingestion_runs":1},"source_sha256":actual_hash},indent=2)+"\n")
    return manifest


def mendoza(source: dict, archive: Path, output: Path, loaded_at: str) -> dict:
    run_id = source["id"]
    facts, source_hashes, rows_read = [], {}, 0
    with zipfile.ZipFile(archive) as bundle:
        for side, member in (("expenditure", "expenditure.csv"), ("revenue", "revenue.csv")):
            payload = bundle.read(member); actual_hash = hashlib.sha256(payload).hexdigest()
            if actual_hash != source["verified_sha256"][side]:
                raise ValueError(f"{side} source hash mismatch: {actual_hash}")
            source_hashes[side] = actual_hash
            reader = csv.reader(io.StringIO(payload.decode("latin1")), delimiter=";")
            header = next(reader); rows = list(reader); rows_read += len(rows)
            enacted_index = 1
            actual_index = next(i for i, value in enumerate(header) if "Acumulado" in value)
            for row_number, row in enumerate(rows, 2):
                label = (row[0] if row else "").strip()
                if not label or label.lower().startswith(("fuente:", "responsable:")):
                    continue
                code = hashlib.sha256(label.encode("utf-8")).hexdigest()[:16]
                upper = label.upper()
                summary = upper.startswith(("TOTAL", "SUBTOTAL", "EROGACIONES ", "RECURSOS ", "DE JURISDICCIÓN"))
                for stage, index in (("revised", enacted_index), ("actual", actual_index)):
                    value = spanish_amount(row[index] if index < len(row) else None)
                    if value is None: continue
                    facts.append({
                        "public_entity_id": "AR:MENDOZA-CIUDAD", "fiscal_year": 2024, "fiscal_period": source["fiscal_period"],
                        "reporting_scope": "standalone_municipality_partial_period", "budget_stage": stage,
                        "budget_side": side, "source_budget_item_type_code": label, "functional_paragraph_code": None,
                        "economic_item_code": code, "amount_local": value, "currency_code": "ARS", "amount_eur": None,
                        "fx_date": None, "is_consolidation_item": False, "is_financing": False, "is_summary_row": summary,
                        "source_row_number": row_number, "source_sheet": member, "source_id": source["id"],
                        "ingestion_run_id": run_id, "quality_flags": ["official_city_open_data", "partial_period_through_april", "anchor_only_not_un_agglomeration", "source_label_hashed_as_code"],
                        "loaded_at": loaded_at,
                    })
    output.mkdir(parents=True, exist_ok=True)
    entities = [{"public_entity_id":"AR:MENDOZA-CIUDAD","entity_name":"Municipalidad de la Ciudad de Mendoza","entity_type":"municipality","country_code_alpha2":"AR","country_code_alpha3":"ARG","national_entity_code":"MENDOZA-CIUDAD","national_entity_code_type":"AR_CITY_PORTAL_CODE","is_eu_capital":False,"is_extra_city":True,"default_currency_code":"ARS","eurostat_city_code":None,"eurostat_geography_name":None,"administrative_region_code":None,"administrative_region_name":"Mendoza","administrative_district_code":None,"administrative_district_name":None,"national_geography_code":"MENDOZA-CIUDAD","national_geography_code_type":"AR_CITY_PORTAL_CODE","valid_from":None,"valid_to":None,"loaded_at":loaded_at}]
    combined_hash = hashlib.sha256("".join(source_hashes[key] for key in sorted(source_hashes)).encode()).hexdigest()
    sources = [{"source_id":source["id"],"public_entity_id":"AR:MENDOZA-CIUDAD","source_type":"official_city_open_data","source_name":"Ejecución Presupuestaria de Gastos y Recursos 2024","source_url":source["landing_url"],"dataset_code":"mendoza-budget-execution-2024-april","archive_file":archive.name,"archive_sha256":combined_hash,"retrieved_at":source["retrieved_at"],"notes":source["fiscal_boundary"] + "; component hashes=" + json.dumps(source_hashes, sort_keys=True),"loaded_at":loaded_at}]
    runs = [{"ingestion_run_id":run_id,"source_id":source["id"],"started_at":loaded_at,"completed_at":loaded_at,"status":"completed","source_vintage":"2024-04","source_sha256":combined_hash,"rows_read":rows_read,"rows_loaded":len(facts),"warning_count":1,"error_message":"Partial period through April; summary rows flagged and must not be added to detail rows."}]
    for name, rows in (("public_entities",entities),("public_entity_sources",sources),("municipal_budget_line_facts",facts),("ingestion_runs",runs)): write_jsonl(output/f"{name}.jsonl",rows)
    manifest={"source_id":source["id"],"source_sha256":combined_hash,"component_sha256":source_hashes,"rows_read":rows_read,"facts":len(facts),"entity_count":1}
    (output/"manifest.json").write_text(json.dumps(manifest,indent=2)+"\n")
    (output/"international_municipal_manifest.json").write_text(json.dumps({"schema_version":"1.0.0","status":"validated","countries":["ARG"],"sources":[source["id"]],"counts":{"public_entities":1,"public_entity_sources":1,"municipal_budget_line_facts":len(facts),"ingestion_runs":1},"source_sha256":combined_hash},indent=2)+"\n")
    return manifest


def guayaquil(source: dict, archive: Path, output: Path, loaded_at: str) -> dict:
    actual_hash = sha256(archive)
    if actual_hash != source["verified_sha256"]:
        raise ValueError(f"Source hash mismatch: {actual_hash}")
    payload = archive.read_bytes().decode("latin1")
    reader = csv.DictReader(io.StringIO(payload), delimiter=";")
    facts, rows_read = [], 0
    columns = (("enacted", "Asignado"), ("revised", "Codificado"), ("committed", "Comprometido"), ("actual", "Devengado"), ("cash", "Pagado"))
    for row_number, row in enumerate(reader, 2):
        rows_read += 1
        code = (row.get("Cuenta") or "").strip()
        label = (row.get("Descripción") or "").strip()
        if not code or not label:
            continue
        for stage, column in columns:
            value = spanish_amount(row.get(column))
            if value is None:
                continue
            facts.append({"public_entity_id":"EC:GUAYAQUIL","fiscal_year":2025,"fiscal_period":source["fiscal_period"],"reporting_scope":"standalone_municipality_partial_period","budget_stage":stage,"budget_side":"expenditure","source_budget_item_type_code":label,"functional_paragraph_code":None,"economic_item_code":code,"amount_local":value,"currency_code":"USD","amount_eur":None,"fx_date":None,"is_consolidation_item":False,"is_financing":False,"is_summary_row":False,"source_row_number":row_number,"source_sheet":"Conjunto-de-datos.csv","source_id":source["id"],"ingestion_run_id":source["id"],"quality_flags":["official_city_transparency_data","partial_period_through_october","anchor_only_not_un_agglomeration"],"loaded_at":loaded_at})
    output.mkdir(parents=True, exist_ok=True)
    entities=[{"public_entity_id":"EC:GUAYAQUIL","entity_name":"Municipalidad de Guayaquil","entity_type":"municipality","country_code_alpha2":"EC","country_code_alpha3":"ECU","national_entity_code":"GUAYAQUIL","national_entity_code_type":"EC_MUNICIPALITY","is_eu_capital":False,"is_extra_city":True,"default_currency_code":"USD","eurostat_city_code":None,"eurostat_geography_name":None,"administrative_region_code":None,"administrative_region_name":"Guayas","administrative_district_code":None,"administrative_district_name":None,"national_geography_code":"GUAYAQUIL","national_geography_code_type":"EC_MUNICIPALITY","valid_from":None,"valid_to":None,"loaded_at":loaded_at}]
    sources=[{"source_id":source["id"],"public_entity_id":"EC:GUAYAQUIL","source_type":"official_city_transparency_data","source_name":"Presupuesto de la Institución, octubre 2025","source_url":source["resource_url"],"dataset_code":"guayaquil-budget-execution-2025-october","archive_file":archive.name,"archive_sha256":actual_hash,"retrieved_at":source["retrieved_at"],"notes":source["fiscal_boundary"] + "; landing=" + source["landing_url"],"loaded_at":loaded_at}]
    runs=[{"ingestion_run_id":source["id"],"source_id":source["id"],"started_at":loaded_at,"completed_at":loaded_at,"status":"completed","source_vintage":"2025-10","source_sha256":actual_hash,"rows_read":rows_read,"rows_loaded":len(facts),"warning_count":1,"error_message":"Partial period through October; anchor municipality only."}]
    for name, rows in (("public_entities",entities),("public_entity_sources",sources),("municipal_budget_line_facts",facts),("ingestion_runs",runs)): write_jsonl(output/f"{name}.jsonl",rows)
    manifest={"source_id":source["id"],"source_sha256":actual_hash,"rows_read":rows_read,"facts":len(facts),"entity_count":1}
    (output/"manifest.json").write_text(json.dumps(manifest,indent=2)+"\n")
    (output/"international_municipal_manifest.json").write_text(json.dumps({"schema_version":"1.0.0","status":"validated","countries":["ECU"],"sources":[source["id"]],"counts":{"public_entities":1,"public_entity_sources":1,"municipal_budget_line_facts":len(facts),"ingestion_runs":1},"source_sha256":actual_hash},indent=2)+"\n")
    return manifest


def panama_city(source: dict, archive: Path, output: Path, loaded_at: str) -> dict:
    actual_hash = sha256(archive)
    if actual_hash != source["verified_sha256"]:
        raise ValueError(f"Source hash mismatch: {actual_hash}")
    sheet = load_workbook(archive, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = [str(value).strip() if value is not None else "" for value in next(rows)]
    indices = {name: header.index(name) for name in ("cta.", "Descripcion", "Presupuesto Ley", "Presupuesto Modificado", "Compromisos / Ejecutado", "Pagado")}
    facts, rows_read = [], 0
    stages = (("enacted", "Presupuesto Ley"), ("revised", "Presupuesto Modificado"), ("actual", "Compromisos / Ejecutado"), ("cash", "Pagado"))
    for row_number, row in enumerate(rows, 2):
        code = str(row[indices["cta."]] or "").strip()
        label = str(row[indices["Descripcion"]] or "").strip()
        if not code or not label:
            continue
        rows_read += 1
        for stage, column in stages:
            raw = row[indices[column]]
            if raw is None:
                continue
            try: value = Decimal(str(raw))
            except InvalidOperation: continue
            if not value: continue
            facts.append({"public_entity_id":"PA:PANAMA-CITY","fiscal_year":source["year"],"fiscal_period":source["fiscal_period"],"reporting_scope":"standalone_municipality","budget_stage":stage,"budget_side":"expenditure","source_budget_item_type_code":label,"functional_paragraph_code":None,"economic_item_code":code,"amount_local":format(value.quantize(Decimal("0.000001")),"f"),"currency_code":source["currency"],"amount_eur":None,"fx_date":None,"is_consolidation_item":False,"is_financing":False,"is_summary_row":False,"source_row_number":row_number,"source_sheet":sheet.title,"source_id":source["id"],"ingestion_run_id":source["id"],"quality_flags":["official_city_transparency_data","anchor_only_not_un_agglomeration"],"loaded_at":loaded_at})
    output.mkdir(parents=True, exist_ok=True)
    entities=[{"public_entity_id":"PA:PANAMA-CITY","entity_name":"Municipio de Panamá","entity_type":"municipality","country_code_alpha2":"PA","country_code_alpha3":"PAN","national_entity_code":"MUPA","national_entity_code_type":"PA_MUNICIPALITY_PORTAL_CODE","is_eu_capital":False,"is_extra_city":True,"default_currency_code":"PAB","eurostat_city_code":None,"eurostat_geography_name":None,"administrative_region_code":None,"administrative_region_name":"Panamá","administrative_district_code":None,"administrative_district_name":None,"national_geography_code":"MUPA","national_geography_code_type":"PA_MUNICIPALITY_PORTAL_CODE","valid_from":None,"valid_to":None,"loaded_at":loaded_at}]
    sources=[{"source_id":source["id"],"public_entity_id":"PA:PANAMA-CITY","source_type":"official_city_transparency_data","source_name":"Informe de Ejecución Presupuestaria Diciembre 2025","source_url":source["resource_url"],"dataset_code":"panama-city-budget-execution-2025-december","archive_file":archive.name,"archive_sha256":actual_hash,"retrieved_at":source["retrieved_at"],"notes":source["fiscal_boundary"] + "; landing=" + source["landing_url"],"loaded_at":loaded_at}]
    runs=[{"ingestion_run_id":source["id"],"source_id":source["id"],"started_at":loaded_at,"completed_at":loaded_at,"status":"completed","source_vintage":"2025-12","source_sha256":actual_hash,"rows_read":rows_read,"rows_loaded":len(facts),"warning_count":1,"error_message":"Anchor municipality only; source labels commitment as Compromisos / Ejecutado."}]
    for name, records in (("public_entities",entities),("public_entity_sources",sources),("municipal_budget_line_facts",facts),("ingestion_runs",runs)): write_jsonl(output/f"{name}.jsonl",records)
    manifest={"source_id":source["id"],"source_sha256":actual_hash,"rows_read":rows_read,"facts":len(facts),"entity_count":1}
    (output/"manifest.json").write_text(json.dumps(manifest,indent=2)+"\n")
    (output/"international_municipal_manifest.json").write_text(json.dumps({"schema_version":"1.0.0","status":"validated","countries":["PAN"],"sources":[source["id"]],"counts":{"public_entities":1,"public_entity_sources":1,"municipal_budget_line_facts":len(facts),"ingestion_runs":1},"source_sha256":actual_hash},indent=2)+"\n")
    return manifest


def brasilia_df(source: dict, archive: Path, output: Path, loaded_at: str) -> dict:
    actual_hash = sha256(archive)
    if actual_hash != source["verified_sha256"]:
        raise ValueError(f"Source hash mismatch: {actual_hash}")
    payload = json.loads(archive.read_text(encoding="utf-8"))
    source_rows = payload.get("items", [])
    if payload.get("hasMore") or len(source_rows) != payload.get("count"):
        raise ValueError("SICONFI response is paginated or incomplete")
    stages = source["stage_mapping"]
    facts=[]
    revenue_columns={"PREVISÃO INICIAL","PREVISÃO ATUALIZADA (a)","Até o Bimestre (c)"}
    expense_columns=set(stages)-revenue_columns
    for row_number,row in enumerate(source_rows,1):
        column=row.get("coluna")
        if column not in stages:
            continue
        value=Decimal(str(row.get("valor") or 0))
        if not value:
            continue
        code=str(row.get("cod_conta") or "").strip()
        label=str(row.get("conta") or "").strip()
        if not code:
            continue
        upper=(code+" "+label).upper()
        summary=code.lower().startswith(("total","receitasexceto","despesasexceto")) or "TOTAL" in upper or " (I)" in upper or " (II)" in upper or " (III)" in upper
        consolidation="INTRA" in upper and "EXCETO" not in upper
        facts.append({"public_entity_id":"BR:53","fiscal_year":source["year"],"fiscal_period":"FY","reporting_scope":"standalone_federal_district","budget_stage":stages[column],"budget_side":"revenue" if column in revenue_columns else "expenditure","source_budget_item_type_code":column,"functional_paragraph_code":None,"economic_item_code":code,"amount_local":format(value.quantize(Decimal("0.000001")),"f"),"currency_code":"BRL","amount_eur":None,"fx_date":None,"is_consolidation_item":consolidation,"is_financing":False,"is_summary_row":summary,"source_row_number":row_number,"source_sheet":"RREO-Anexo 01","source_id":source["id"],"ingestion_run_id":source["id"],"quality_flags":["official_siconfi_rreo","federal_district_anchor_only_not_un_agglomeration"],"loaded_at":loaded_at})
    output.mkdir(parents=True,exist_ok=True)
    entities=[{"public_entity_id":"BR:53","entity_name":"Governo do Distrito Federal","entity_type":"federal_district","country_code_alpha2":"BR","country_code_alpha3":"BRA","national_entity_code":"53","national_entity_code_type":"BR_IBGE_STATE_CODE","is_eu_capital":False,"is_extra_city":True,"default_currency_code":"BRL","eurostat_city_code":None,"eurostat_geography_name":None,"administrative_region_code":"53","administrative_region_name":"Distrito Federal","administrative_district_code":None,"administrative_district_name":None,"national_geography_code":"53","national_geography_code_type":"BR_IBGE_STATE_CODE","valid_from":None,"valid_to":None,"loaded_at":loaded_at}]
    sources=[{"source_id":source["id"],"public_entity_id":"BR:53","source_type":"official_national_fiscal_api","source_name":"SICONFI RREO Anexo 01, 6º bimestre 2025","source_url":source["resource_url"],"dataset_code":"siconfi-rreo-anexo-01","archive_file":archive.name,"archive_sha256":actual_hash,"retrieved_at":source["retrieved_at"],"notes":source["fiscal_boundary"] + "; landing=" + source["landing_url"],"loaded_at":loaded_at}]
    runs=[{"ingestion_run_id":source["id"],"source_id":source["id"],"started_at":loaded_at,"completed_at":loaded_at,"status":"completed","source_vintage":"2025-B6","source_sha256":actual_hash,"rows_read":len(source_rows),"rows_loaded":len(facts),"warning_count":1,"error_message":"Distrito Federal is a federal-district anchor, not a consolidated UN urban agglomeration."}]
    for name,records in (("public_entities",entities),("public_entity_sources",sources),("municipal_budget_line_facts",facts),("ingestion_runs",runs)): write_jsonl(output/f"{name}.jsonl",records)
    manifest={"source_id":source["id"],"source_sha256":actual_hash,"rows_read":len(source_rows),"facts":len(facts),"entity_count":1}
    (output/"manifest.json").write_text(json.dumps(manifest,indent=2)+"\n")
    (output/"international_municipal_manifest.json").write_text(json.dumps({"schema_version":"1.0.0","status":"validated","countries":["BRA"],"sources":[source["id"]],"counts":{"public_entities":1,"public_entity_sources":1,"municipal_budget_line_facts":len(facts),"ingestion_runs":1},"source_sha256":actual_hash},indent=2)+"\n")
    return manifest


def bogota_cuipo(source: dict, archive: Path, output: Path, loaded_at: str) -> dict:
    facts=[]; rows_read=0; component_hashes={}
    dimensions=("cuenta","cod_vigencia_del_gasto","cod_seccion_presupuestal","cod_sector","cod_programatico_mga","bpin","cod_fuentes_financiacion","cod_situacion_de_fondos","cod_politica_publica","cod_terceros")
    with zipfile.ZipFile(archive) as bundle:
        for member, measures in (("programming.csv",(("enacted","apropiacion_inicial"),("revised","apropiacion_definitiva"))), ("execution.csv",(("committed","compromisos"),("actual","obligaciones"),("cash","pagos")))):
            payload=bundle.read(member); component_hashes[member]=hashlib.sha256(payload).hexdigest()
            reader=csv.DictReader(io.StringIO(payload.decode("utf-8-sig")))
            for row_number,row in enumerate(reader,2):
                rows_read+=1
                if row.get("codigo_entidad") != source["entity_code"] or row.get("periodo") != "20251201":
                    raise ValueError(f"Unexpected CUIPO entity/period in {member} row {row_number}")
                raw_parts=[(row.get(key) or "0").strip() for key in dimensions]
                composite="|".join(raw_parts)
                label=(row.get("nombre_cuenta") or "").strip()
                for stage,column in measures:
                    raw=(row.get(column) or "").strip()
                    if not raw: continue
                    try: value=Decimal(raw)
                    except InvalidOperation: continue
                    if not value: continue
                    facts.append({"public_entity_id":"CO:210111001","fiscal_year":2025,"fiscal_period":"FY","reporting_scope":"standalone_capital_district","budget_stage":stage,"budget_side":"expenditure","source_budget_item_type_code":label,"functional_paragraph_code":(row.get("cod_programatico_mga") or None),"economic_item_code":composite,"amount_local":format(value.quantize(Decimal("0.000001")),"f"),"currency_code":"COP","amount_eur":None,"fx_date":None,"is_consolidation_item":False,"is_financing":False,"is_summary_row":False,"source_row_number":row_number,"source_sheet":member,"source_id":source["id"],"ingestion_run_id":source["id"],"quality_flags":["official_colombia_cuipo","anchor_only_not_un_agglomeration","composite_classification_key"],"loaded_at":loaded_at})
    output.mkdir(parents=True,exist_ok=True)
    combined_hash=hashlib.sha256("".join(component_hashes[k] for k in sorted(component_hashes)).encode()).hexdigest()
    entities=[{"public_entity_id":"CO:210111001","entity_name":"Bogotá D.C.","entity_type":"capital_district","country_code_alpha2":"CO","country_code_alpha3":"COL","national_entity_code":"210111001","national_entity_code_type":"CO_CUIPO_ENTITY_CODE","is_eu_capital":False,"is_extra_city":True,"default_currency_code":"COP","eurostat_city_code":None,"eurostat_geography_name":None,"administrative_region_code":None,"administrative_region_name":"Bogotá D.C.","administrative_district_code":None,"administrative_district_name":None,"national_geography_code":"210111001","national_geography_code_type":"CO_CUIPO_ENTITY_CODE","valid_from":None,"valid_to":None,"loaded_at":loaded_at}]
    sources=[{"source_id":source["id"],"public_entity_id":"CO:210111001","source_type":"official_national_open_data_api","source_name":"OVCF CUIPO Programación y Ejecución de Gastos 2025","source_url":source["landing_url"],"dataset_code":"d9mu-h6ar+4f7r-epif","archive_file":archive.name,"archive_sha256":combined_hash,"retrieved_at":source["retrieved_at"],"notes":source["fiscal_boundary"] + "; component_sha256=" + json.dumps(component_hashes,sort_keys=True) + "; resource_urls=" + json.dumps(source["resource_urls"],sort_keys=True),"loaded_at":loaded_at}]
    runs=[{"ingestion_run_id":source["id"],"source_id":source["id"],"started_at":loaded_at,"completed_at":loaded_at,"status":"completed","source_vintage":"2025-Q4","source_sha256":combined_hash,"rows_read":rows_read,"rows_loaded":len(facts),"warning_count":1,"error_message":"Bogotá D.C. anchor jurisdiction only; CUIPO dimensional keys preserved as a composite classification code."}]
    for name,records in (("public_entities",entities),("public_entity_sources",sources),("municipal_budget_line_facts",facts),("ingestion_runs",runs)): write_jsonl(output/f"{name}.jsonl",records)
    manifest={"source_id":source["id"],"source_sha256":combined_hash,"component_sha256":component_hashes,"rows_read":rows_read,"facts":len(facts),"entity_count":1}
    (output/"manifest.json").write_text(json.dumps(manifest,indent=2)+"\n")
    (output/"international_municipal_manifest.json").write_text(json.dumps({"schema_version":"1.0.0","status":"validated","countries":["COL"],"sources":[source["id"]],"counts":{"public_entities":1,"public_entity_sources":1,"municipal_budget_line_facts":len(facts),"ingestion_runs":1},"source_sha256":combined_hash},indent=2)+"\n")
    return manifest


def colombia_four_cities(source: dict, archive: Path, output: Path, loaded_at: str) -> dict:
    names={"210105001":"Distrito Especial de Ciencia, Tecnología e Innovación de Medellín","210176001":"Santiago de Cali","210108001":"Barranquilla, Distrito Especial, Industrial y Portuario","210113001":"Cartagena de Indias, Distrito Turístico y Cultural"}
    regions={"210105001":"Antioquia","210176001":"Valle del Cauca","210108001":"Atlántico","210113001":"Bolívar"}
    dimensions=("cuenta","cod_vigencia_del_gasto","cod_seccion_presupuestal","cod_sector","cod_programatico_mga","bpin","cod_fuentes_financiacion","cod_situacion_de_fondos","cod_politica_publica","cod_terceros")
    facts=[]; rows_read=0; hashes={}; seen=set()
    with zipfile.ZipFile(archive) as bundle:
        for member in sorted(bundle.namelist()):
            payload=bundle.read(member); hashes[member]=hashlib.sha256(payload).hexdigest()
            code=member.split("_",1)[0]; kind="programming" if "programming" in member else "execution"
            measures=(("enacted","apropiacion_inicial"),("revised","apropiacion_definitiva")) if kind=="programming" else (("committed","compromisos"),("actual","obligaciones"),("cash","pagos"))
            for row_number,row in enumerate(csv.DictReader(io.StringIO(payload.decode("utf-8-sig"))),2):
                rows_read+=1; seen.add(code)
                if row.get("codigo_entidad") != code or row.get("periodo") != "20251201": raise ValueError(f"Unexpected CUIPO row in {member}:{row_number}")
                composite="|".join((row.get(key) or "0").strip() for key in dimensions)
                for stage,column in measures:
                    raw=(row.get(column) or "").strip()
                    if not raw: continue
                    try: value=Decimal(raw)
                    except InvalidOperation: continue
                    if not value: continue
                    facts.append({"public_entity_id":"CO:"+code,"fiscal_year":2025,"fiscal_period":"FY","reporting_scope":"standalone_municipality","budget_stage":stage,"budget_side":"expenditure","source_budget_item_type_code":row.get("nombre_cuenta") or None,"functional_paragraph_code":row.get("cod_programatico_mga") or None,"economic_item_code":composite,"amount_local":format(value.quantize(Decimal("0.000001")),"f"),"currency_code":"COP","amount_eur":None,"fx_date":None,"is_consolidation_item":False,"is_financing":False,"is_summary_row":False,"source_row_number":row_number,"source_sheet":member,"source_id":source["id"],"ingestion_run_id":source["id"],"quality_flags":["official_colombia_cuipo","anchor_only_not_un_agglomeration","composite_classification_key"],"loaded_at":loaded_at})
    if seen != set(source["entity_codes"]): raise ValueError(f"Missing CUIPO entities: {set(source['entity_codes'])-seen}")
    output.mkdir(parents=True,exist_ok=True); combined=hashlib.sha256("".join(hashes[k] for k in sorted(hashes)).encode()).hexdigest()
    entities=[]; sources=[]
    for code in source["entity_codes"]:
        entities.append({"public_entity_id":"CO:"+code,"entity_name":names[code],"entity_type":"municipality","country_code_alpha2":"CO","country_code_alpha3":"COL","national_entity_code":code,"national_entity_code_type":"CO_CUIPO_ENTITY_CODE","is_eu_capital":False,"is_extra_city":True,"default_currency_code":"COP","eurostat_city_code":None,"eurostat_geography_name":None,"administrative_region_code":None,"administrative_region_name":regions[code],"administrative_district_code":None,"administrative_district_name":None,"national_geography_code":code,"national_geography_code_type":"CO_CUIPO_ENTITY_CODE","valid_from":None,"valid_to":None,"loaded_at":loaded_at})
        sources.append({"source_id":source["id"],"public_entity_id":"CO:"+code,"source_type":"official_national_open_data_api","source_name":"OVCF CUIPO Programación y Ejecución de Gastos 2025","source_url":source["landing_url"],"dataset_code":"d9mu-h6ar+4f7r-epif","archive_file":archive.name,"archive_sha256":combined,"retrieved_at":source["retrieved_at"],"notes":source["fiscal_boundary"]+"; component_sha256="+json.dumps(hashes,sort_keys=True),"loaded_at":loaded_at})
    runs=[{"ingestion_run_id":source["id"],"source_id":source["id"],"started_at":loaded_at,"completed_at":loaded_at,"status":"completed","source_vintage":"2025-Q4","source_sha256":combined,"rows_read":rows_read,"rows_loaded":len(facts),"warning_count":4,"error_message":"Four anchor jurisdictions only; CUIPO dimensional keys preserved as composite classification codes."}]
    for name,records in (("public_entities",entities),("public_entity_sources",sources),("municipal_budget_line_facts",facts),("ingestion_runs",runs)): write_jsonl(output/f"{name}.jsonl",records)
    manifest={"source_id":source["id"],"source_sha256":combined,"component_sha256":hashes,"rows_read":rows_read,"facts":len(facts),"entity_count":len(entities)}
    (output/"manifest.json").write_text(json.dumps(manifest,indent=2)+"\n"); (output/"international_municipal_manifest.json").write_text(json.dumps({"schema_version":"1.0.0","status":"validated","countries":["COL"],"sources":[source["id"]],"counts":{"public_entities":len(entities),"public_entity_sources":len(sources),"municipal_budget_line_facts":len(facts),"ingestion_runs":1},"source_sha256":combined},indent=2)+"\n")
    return manifest


def san_jose_cr(source: dict, archive: Path, output: Path, loaded_at: str) -> dict:
    actual_hash=sha256(archive)
    if actual_hash != source["verified_sha256"]: raise ValueError(f"Source hash mismatch: {actual_hash}")
    sheet=load_workbook(archive,read_only=True,data_only=True).active
    facts=[]; rows_read=0
    for row_number,row in enumerate(sheet.iter_rows(min_row=8,values_only=True),8):
        code=str(row[0] or "").strip(); label=str(row[1] or "").strip()
        if not code or not label: continue
        rows_read+=1
        raw=row[8] if len(row)>8 else None
        if raw is None: continue
        try: value=Decimal(str(raw))
        except InvalidOperation: continue
        if not value: continue
        facts.append({"public_entity_id":"CR:SAN-JOSE","fiscal_year":2025,"fiscal_period":"2025-06","reporting_scope":"standalone_municipality_partial_period","budget_stage":"actual","budget_side":"expenditure","source_budget_item_type_code":label,"functional_paragraph_code":None,"economic_item_code":code,"amount_local":format(value.quantize(Decimal("0.000001")),"f"),"currency_code":"CRC","amount_eur":None,"fx_date":None,"is_consolidation_item":False,"is_financing":False,"is_summary_row":False,"source_row_number":row_number,"source_sheet":sheet.title,"source_id":source["id"],"ingestion_run_id":source["id"],"quality_flags":["official_city_budget_workbook","partial_period_through_june","anchor_only_not_un_agglomeration","tls_chain_invalid_source_hash_pinned"],"loaded_at":loaded_at})
    output.mkdir(parents=True,exist_ok=True)
    entities=[{"public_entity_id":"CR:SAN-JOSE","entity_name":"Municipalidad de San José","entity_type":"municipality","country_code_alpha2":"CR","country_code_alpha3":"CRI","national_entity_code":"SAN-JOSE","national_entity_code_type":"CR_MUNICIPALITY_NAME_CODE","is_eu_capital":False,"is_extra_city":True,"default_currency_code":"CRC","eurostat_city_code":None,"eurostat_geography_name":None,"administrative_region_code":None,"administrative_region_name":"San José","administrative_district_code":None,"administrative_district_name":None,"national_geography_code":"SAN-JOSE","national_geography_code_type":"CR_MUNICIPALITY_NAME_CODE","valid_from":None,"valid_to":None,"loaded_at":loaded_at}]
    sources=[{"source_id":source["id"],"public_entity_id":"CR:SAN-JOSE","source_type":"official_city_budget_workbook","source_name":"Egresos mensuales, 1 enero al 30 junio 2025","source_url":source["resource_url"],"dataset_code":"ejecutado-al-30-de-junio-2025","archive_file":archive.name,"archive_sha256":actual_hash,"retrieved_at":source["retrieved_at"],"notes":source["fiscal_boundary"]+"; official host served an incomplete TLS chain; exact workbook SHA-256 pinned before production ingestion; landing="+source["landing_url"],"loaded_at":loaded_at}]
    runs=[{"ingestion_run_id":source["id"],"source_id":source["id"],"started_at":loaded_at,"completed_at":loaded_at,"status":"completed","source_vintage":"2025-06","source_sha256":actual_hash,"rows_read":rows_read,"rows_loaded":len(facts),"warning_count":2,"error_message":"Partial period through June; official host TLS chain invalid, content accepted only after exact SHA-256 match."}]
    for name,records in (("public_entities",entities),("public_entity_sources",sources),("municipal_budget_line_facts",facts),("ingestion_runs",runs)): write_jsonl(output/f"{name}.jsonl",records)
    manifest={"source_id":source["id"],"source_sha256":actual_hash,"rows_read":rows_read,"facts":len(facts),"entity_count":1}
    (output/"manifest.json").write_text(json.dumps(manifest,indent=2)+"\n"); (output/"international_municipal_manifest.json").write_text(json.dumps({"schema_version":"1.0.0","status":"validated","countries":["CRI"],"sources":[source["id"]],"counts":{"public_entities":1,"public_entity_sources":1,"municipal_budget_line_facts":len(facts),"ingestion_runs":1},"source_sha256":actual_hash},indent=2)+"\n")
    return manifest


def montevideo(source: dict, archive: Path, output: Path, loaded_at: str) -> dict:
    expected = source.get("verified_sha256")
    actual = sha256(archive)
    if expected and actual != expected:
        raise ValueError(f"Source hash mismatch: {actual}")
    run_id = f"{source['id']}-{actual[:12]}"
    facts: list[dict] = []
    raw_rows = 0
    for side, prefix, code_columns, name_columns, measures in (
        ("revenue", "ingresos_", ("cuenta",), ("descripcion_cuenta",), (("enacted", "monto_autorizado"), ("actual", "monto_recaudado"))),
        ("expenditure", "egresos_", ("area_funcional", "pos_presupuestaria"), ("descripcion_area_funcional", "descripcion_pos_presupuestaria"), (("enacted", "monto_autorizado"), ("actual", "monto_incorporado"))),
    ):
        for member, reader in rows_from_zip(archive, prefix):
            for row_number, row in enumerate(reader, 2):
                raw_rows += 1
                code = ":".join((row.get(key) or "UNSPECIFIED").strip() for key in code_columns)
                label = " · ".join((row.get(key) or "").strip() for key in name_columns if (row.get(key) or "").strip())
                for stage, column in measures:
                    value = amount(row.get(column))
                    if value is None:
                        continue
                    facts.append({
                        "public_entity_id": "UY:MONTEVIDEO", "fiscal_year": source["year"], "fiscal_period": "FY",
                        "reporting_scope": "standalone_subnational_government", "budget_stage": stage,
                        "budget_side": side, "source_budget_item_type_code": label or None,
                        "functional_paragraph_code": (row.get("area_funcional") or None),
                        "economic_item_code": code, "amount_local": value,
                        "currency_code": source["currency"], "amount_eur": None, "fx_date": None,
                        "is_consolidation_item": False,
                        "is_financing": False, "is_summary_row": False, "source_row_number": row_number,
                        "source_sheet": member, "source_id": source["id"], "ingestion_run_id": run_id,
                        "quality_flags": ["official_city_open_data", "anchor_only_not_un_agglomeration"],
                        "loaded_at": loaded_at,
                    })
    output.mkdir(parents=True, exist_ok=True)
    entity = [{
        "public_entity_id": "UY:MONTEVIDEO", "entity_name": "Intendencia de Montevideo",
        "entity_type": "departmental_government", "country_code_alpha2": "UY", "country_code_alpha3": "URY",
        "national_entity_code": source["entity_code"], "national_entity_code_type": "UY_DEPARTMENTAL_GOVERNMENT",
        "is_eu_capital": False, "is_extra_city": True, "default_currency_code": source["currency"],
        "eurostat_city_code": None, "eurostat_geography_name": None,
        "administrative_region_code": None, "administrative_region_name": "Montevideo",
        "administrative_district_code": None, "administrative_district_name": None,
        "national_geography_code": source["entity_code"], "national_geography_code_type": "UY_DEPARTMENTAL_GOVERNMENT",
        "valid_from": None, "valid_to": None, "loaded_at": loaded_at,
    }]
    provenance = [{
        "source_id": source["id"], "public_entity_id": "UY:MONTEVIDEO", "source_type": "official_city_open_data",
        "source_name": "Balance de ejecución presupuestal", "source_url": source["resource_url"],
        "dataset_code": "balance-de-ejecucion-presupuestal", "archive_file": archive.name,
        "archive_sha256": actual, "retrieved_at": source["retrieved_at"], "loaded_at": loaded_at,
        "notes": f"{source['fiscal_boundary']}; landing={source['landing_url']}; metadata={source['metadata_url']}",
    }]
    run = [{
        "ingestion_run_id": run_id, "source_id": source["id"], "started_at": loaded_at,
        "completed_at": loaded_at, "status": "completed", "source_vintage": str(source["year"]),
        "source_sha256": actual, "rows_read": raw_rows, "rows_loaded": len(facts),
        "warning_count": 0, "error_message": None,
    }]
    write_jsonl(output / "public_entities.jsonl", entity)
    write_jsonl(output / "public_entity_sources.jsonl", provenance)
    write_jsonl(output / "municipal_budget_line_facts.jsonl", facts)
    write_jsonl(output / "ingestion_runs.jsonl", run)
    manifest = {"source_id": source["id"], "source_sha256": actual, "rows_read": raw_rows, "facts": len(facts), "entity_count": 1}
    (output / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    (output / "international_municipal_manifest.json").write_text(json.dumps({
        "schema_version": "1.0.0", "status": "validated", "countries": ["URY"],
        "sources": [source["id"]], "counts": {"public_entities": 1, "public_entity_sources": 1,
        "municipal_budget_line_facts": len(facts), "ingestion_runs": 1}, "source_sha256": actual,
    }, indent=2) + "\n", encoding="utf-8")
    return manifest


def sao_paulo(source: dict, archive: Path, output: Path, loaded_at: str) -> dict:
    actual_hash = sha256(archive)
    if actual_hash != source["verified_sha256"]:
        raise ValueError(f"Source hash mismatch: {actual_hash}")
    reader = csv.DictReader(io.StringIO(archive.read_text(encoding="latin-1")), delimiter=";")
    required = {
        "Cd_AnoExecucao", "Cd_Exercicio", "Cd_Dotacao_Id", "Cd_Orgao", "Cd_Unidade",
        "Cd_Funcao", "Cd_SubFuncao", "Cd_Programa", "ProjetoAtividade", "Cd_Despesa",
        "Cd_Fonte", "COD_EX_FONT_REC", "COD_DSTN_REC", "COD_VINC_REC_PMSP",
        "Vl_Orcado_Ano", "Vl_Orcado_Atualizado", "Vl_EmpenhadoLiquido", "Vl_Liquidado", "Vl_Pago",
    }
    if not reader.fieldnames or not required.issubset(reader.fieldnames):
        raise ValueError(f"Unexpected São Paulo columns: {reader.fieldnames}")
    stages = (
        ("enacted", "Vl_Orcado_Ano"), ("revised", "Vl_Orcado_Atualizado"),
        ("committed", "Vl_EmpenhadoLiquido"), ("actual", "Vl_Liquidado"), ("cash", "Vl_Pago"),
    )
    facts, rows_read = [], 0
    for row_number, row in enumerate(reader, 2):
        if row.get("Cd_AnoExecucao") != "2025" or row.get("Cd_Exercicio") != "2025":
            raise ValueError(f"Unexpected São Paulo fiscal year at row {row_number}")
        dotacao = (row.get("Cd_Dotacao_Id") or "").strip()
        expense = (row.get("Cd_Despesa") or "").strip()
        if not dotacao or not expense:
            raise ValueError(f"Missing appropriation/expense code at row {row_number}")
        rows_read += 1
        composite = "|".join((row.get(key) or "").strip() for key in (
            "Cd_Orgao", "Cd_Unidade", "Cd_Funcao", "Cd_SubFuncao", "Cd_Programa",
            "ProjetoAtividade", "Cd_Despesa", "Cd_Fonte", "COD_EX_FONT_REC",
            "COD_DSTN_REC", "COD_VINC_REC_PMSP", "COD_TIP_CRED_ORCM", "COD_RDZD_FONT_REC",
        ))
        label = " / ".join(filter(None, ((row.get(key) or "").strip() for key in (
            "Ds_Orgao", "Ds_Unidade", "Ds_Programa", "Ds_Projeto_Atividade", "Ds_Despesa",
        ))))
        for stage, column in stages:
            value = spanish_amount(row.get(column))
            if value is None:
                continue
            facts.append({
                "public_entity_id":"BR:3550308", "fiscal_year":2025, "fiscal_period":"FY",
                "reporting_scope":"standalone_municipality", "budget_stage":stage,
                "budget_side":"expenditure", "source_budget_item_type_code":label,
                "functional_paragraph_code":"|".join(((row.get("Cd_Funcao") or "").strip(), (row.get("Cd_SubFuncao") or "").strip(), (row.get("Cd_Programa") or "").strip(), (row.get("ProjetoAtividade") or "").strip())),
                "economic_item_code":dotacao + "|" + composite, "amount_local":value,
                "currency_code":"BRL", "amount_eur":None, "fx_date":None,
                "is_consolidation_item":False, "is_financing":False, "is_summary_row":False,
                "source_row_number":row_number, "source_sheet":"basedadosexecucao_1225.csv",
                "source_id":source["id"], "ingestion_run_id":source["id"],
                "quality_flags":["official_city_open_data", "maximum_detail_appropriation_execution", "anchor_only_not_un_agglomeration", "replaces_siconfi_rreo_aggregate_2025"],
                "loaded_at":loaded_at,
            })
    output.mkdir(parents=True, exist_ok=True)
    entities=[{"public_entity_id":"BR:3550308","entity_name":"Município de São Paulo","entity_type":"municipality","country_code_alpha2":"BR","country_code_alpha3":"BRA","national_entity_code":"3550308","national_entity_code_type":"BR_IBGE_MUNICIPALITY_CODE","is_eu_capital":False,"is_extra_city":True,"default_currency_code":"BRL","eurostat_city_code":None,"eurostat_geography_name":None,"administrative_region_code":"35","administrative_region_name":"São Paulo","administrative_district_code":None,"administrative_district_name":None,"national_geography_code":"3550308","national_geography_code_type":"BR_IBGE_MUNICIPALITY_CODE","valid_from":None,"valid_to":None,"loaded_at":loaded_at}]
    sources=[{"source_id":source["id"],"public_entity_id":"BR:3550308","source_type":"official_city_open_data","source_name":"Base de Dados da Execução Orçamentária 2025, posição 31/12/2025","source_url":source["resource_url"],"dataset_code":"basedadosexecucao-1225","archive_file":archive.name,"archive_sha256":actual_hash,"retrieved_at":source["retrieved_at"],"notes":source["fiscal_boundary"]+"; landing="+source["landing_url"]+"; maximum-detail appropriation rows replace overlapping 2025 SICONFI RREO aggregate facts for this entity","loaded_at":loaded_at}]
    runs=[{"ingestion_run_id":source["id"],"source_id":source["id"],"started_at":loaded_at,"completed_at":loaded_at,"status":"completed","source_vintage":"2025-12-31","source_sha256":actual_hash,"rows_read":rows_read,"rows_loaded":len(facts),"warning_count":1,"error_message":"Anchor municipality only; overlapping 2025 SICONFI RREO aggregate facts must be removed after verified load."}]
    for name,records in (("public_entities",entities),("public_entity_sources",sources),("municipal_budget_line_facts",facts),("ingestion_runs",runs)): write_jsonl(output/f"{name}.jsonl",records)
    manifest={"source_id":source["id"],"source_sha256":actual_hash,"rows_read":rows_read,"facts":len(facts),"entity_count":1,"replacement":{"public_entity_id":"BR:3550308","fiscal_year":2025,"source_id":"br-siconfi-rreo-2025","expected_rows":288}}
    (output/"manifest.json").write_text(json.dumps(manifest,indent=2)+"\n")
    (output/"international_municipal_manifest.json").write_text(json.dumps({"schema_version":"1.0.0","status":"validated","countries":["BRA"],"sources":[source["id"]],"counts":{"public_entities":1,"public_entity_sources":1,"municipal_budget_line_facts":len(facts),"ingestion_runs":1},"source_sha256":actual_hash},indent=2)+"\n")
    return manifest


def lima_mef(source: dict, archive: Path, output: Path, loaded_at: str) -> dict:
    actual_hash = sha256(archive)
    if actual_hash != source["verified_sha256"]:
        raise ValueError(f"Source hash mismatch: {actual_hash}")
    reader = csv.DictReader(io.StringIO(archive.read_text(encoding="utf-8-sig")))
    dimensions = (
        "SEC_FUNC", "PROGRAMA_PPTO", "TIPO_ACT_PROY", "PRODUCTO_PROYECTO",
        "ACTIVIDAD_ACCION_OBRA", "FUNCION", "DIVISION_FUNCIONAL", "GRUPO_FUNCIONAL",
        "META", "FINALIDAD", "DEPARTAMENTO_META", "FUENTE_FINANCIAMIENTO", "RUBRO",
        "TIPO_RECURSO", "CATEGORIA_GASTO", "TIPO_TRANSACCION", "GENERICA",
        "SUBGENERICA", "SUBGENERICA_DET", "ESPECIFICA", "ESPECIFICA_DET",
    )
    required = {
        "ANO_EJE", "MES_EJE", "NIVEL_GOBIERNO", "SEC_EJEC", "EJECUTORA_NOMBRE",
        *dimensions, *source["stage_mapping"],
    }
    if not reader.fieldnames or not required.issubset(reader.fieldnames):
        raise ValueError(f"Unexpected Lima MEF columns: {reader.fieldnames}")
    facts, rows_read = [], 0
    label_fields = (
        "PROGRAMA_PPTO_NOMBRE", "PRODUCTO_PROYECTO_NOMBRE", "ACTIVIDAD_ACCION_OBRA_NOMBRE",
        "META_NOMBRE", "ESPECIFICA_DET_NOMBRE",
    )
    for row_number, row in enumerate(reader, 2):
        if row["ANO_EJE"] != "2024" or row["NIVEL_GOBIERNO"] != "M" or row["SEC_EJEC"] != "301250":
            raise ValueError(f"Unexpected Lima entity/year boundary at row {row_number}")
        if row["EJECUTORA_NOMBRE"].strip() != "MUNICIPALIDAD METROPOLITANA DE LIMA":
            raise ValueError(f"Unexpected Lima legal entity name at row {row_number}")
        month = row["MES_EJE"].strip()
        if not month.isdigit() or not 0 <= int(month) <= 12:
            raise ValueError(f"Unexpected execution month at row {row_number}: {month!r}")
        rows_read += 1
        composite = "|".join(row.get(key, "").strip() for key in dimensions)
        label = " / ".join(filter(None, (row.get(key, "").strip() for key in label_fields)))
        fiscal_period = "FY" if month == "0" else f"2024-{int(month):02d}"
        for column, stage in source["stage_mapping"].items():
            value = spanish_amount(row.get(column))
            if value is None:
                continue
            facts.append({
                "public_entity_id":"PE:301250", "fiscal_year":2024,
                "fiscal_period":fiscal_period, "reporting_scope":"standalone_metropolitan_municipality",
                "budget_stage":stage, "budget_side":"expenditure",
                "source_budget_item_type_code":label,
                "functional_paragraph_code":"|".join(row.get(key, "").strip() for key in (
                    "FUNCION", "DIVISION_FUNCIONAL", "GRUPO_FUNCIONAL", "PROGRAMA_PPTO",
                    "PRODUCTO_PROYECTO", "ACTIVIDAD_ACCION_OBRA", "META", "FINALIDAD",
                )),
                "economic_item_code":composite, "amount_local":value, "currency_code":"PEN",
                "amount_eur":None, "fx_date":None, "is_consolidation_item":False,
                "is_financing":False, "is_summary_row":False, "source_row_number":row_number,
                "source_sheet":"2024-Gasto.csv filtered SEC_EJEC=301250", "source_id":source["id"],
                "ingestion_run_id":source["id"],
                "quality_flags":["official_national_open_data", "maximum_detail_mef_execution", "exact_sec_ejec_filter", "anchor_only_not_un_agglomeration"],
                "loaded_at":loaded_at,
            })
    output.mkdir(parents=True, exist_ok=True)
    entities=[{"public_entity_id":"PE:301250","entity_name":"Municipalidad Metropolitana de Lima","entity_type":"metropolitan_municipality","country_code_alpha2":"PE","country_code_alpha3":"PER","national_entity_code":"301250","national_entity_code_type":"PE_MEF_SEC_EJEC","is_eu_capital":False,"is_extra_city":True,"default_currency_code":"PEN","eurostat_city_code":None,"eurostat_geography_name":None,"administrative_region_code":"15","administrative_region_name":"Lima","administrative_district_code":"150101","administrative_district_name":"Lima","national_geography_code":"301250","national_geography_code_type":"PE_MEF_SEC_EJEC","valid_from":None,"valid_to":None,"loaded_at":loaded_at}]
    sources=[{"source_id":source["id"],"public_entity_id":"PE:301250","source_type":"official_national_open_data","source_name":"MEF 2024 Presupuesto y Ejecución de Gasto, SEC_EJEC 301250","source_url":source["resource_url"],"dataset_code":source["resource_id"],"archive_file":archive.name,"archive_sha256":actual_hash,"retrieved_at":source["retrieved_at"],"notes":source["fiscal_boundary"]+"; landing="+source["landing_url"]+"; full official ZIP="+source["official_full_zip_url"],"loaded_at":loaded_at}]
    runs=[{"ingestion_run_id":source["id"],"source_id":source["id"],"started_at":loaded_at,"completed_at":loaded_at,"status":"completed","source_vintage":"2024","source_sha256":actual_hash,"rows_read":rows_read,"rows_loaded":len(facts),"warning_count":1,"error_message":"Anchor municipality only; monthly execution rows and FY opening/modified appropriation rows retained at source granularity."}]
    for name,records in (("public_entities",entities),("public_entity_sources",sources),("municipal_budget_line_facts",facts),("ingestion_runs",runs)): write_jsonl(output/f"{name}.jsonl",records)
    manifest={"source_id":source["id"],"source_sha256":actual_hash,"rows_read":rows_read,"facts":len(facts),"entity_count":1,"replacement":{"public_entity_id":"PE:301250","fiscal_year":2024,"source_id":"per-municipal-expansion","expected_rows":81}}
    (output/"manifest.json").write_text(json.dumps(manifest,indent=2)+"\n")
    (output/"international_municipal_manifest.json").write_text(json.dumps({"schema_version":"1.0.0","status":"validated","countries":["PER"],"sources":[source["id"]],"counts":{"public_entities":1,"public_entity_sources":1,"municipal_budget_line_facts":len(facts),"ingestion_runs":1},"source_sha256":actual_hash},indent=2)+"\n")
    return manifest


def rio_de_janeiro(source: dict, archive: Path, output: Path, loaded_at: str) -> dict:
    actual_hash = sha256(archive)
    if actual_hash != source["verified_sha256"]:
        raise ValueError(f"Source hash mismatch: {actual_hash}")
    reader = csv.DictReader(io.StringIO(archive.read_text(encoding="latin-1")), delimiter=";")
    dimensions = (
        "Orgao", "Unidade_Orcamentaria", "Codigo_Funcao", "Codigo_Subfuncao",
        "Codigo_Programa", "Codigo_Acao", "Codigo_Programa_Trabalho", "Fonte_Recurso",
        "Fonte_Recurso_Detalhada", "Categoria_Economica", "Codigo_Grupo",
        "Codigo_Modalidade_Aplicacao", "Codigo_Elemento", "Codigo_Natureza_Despesa",
        "Codigo_Item_Patrimonial",
    )
    required = {"Exercicio", *dimensions, *source["stage_mapping"]}
    if not reader.fieldnames or not required.issubset(reader.fieldnames):
        raise ValueError(f"Unexpected Rio columns: {reader.fieldnames}")
    facts, rows_read = [], 0
    label_fields = (
        "Descricao_Orgao", "Descricao_Unidade_Orcamentaria", "Descricao_Programa",
        "Descricao_Acao", "Descricao_Programa_Trabalho", "Descricao_Natureza Despesa",
        "Descricao_Item Patrimonial",
    )
    for row_number, row in enumerate(reader, 2):
        if row["Exercicio"].strip() != "2025":
            raise ValueError(f"Unexpected Rio fiscal year at row {row_number}")
        if not row["Codigo_Programa_Trabalho"].strip() or not row["Codigo_Natureza_Despesa"].strip():
            raise ValueError(f"Missing Rio program/nature code at row {row_number}")
        rows_read += 1
        composite = "|".join(row.get(key, "").strip() for key in dimensions)
        label = " / ".join(filter(None, (row.get(key, "").strip() for key in label_fields)))
        for column, stage in source["stage_mapping"].items():
            value = spanish_amount(row.get(column))
            if value is None:
                continue
            facts.append({
                "public_entity_id":"BR:3304557", "fiscal_year":2025, "fiscal_period":"FY",
                "reporting_scope":"standalone_municipality", "budget_stage":stage,
                "budget_side":"expenditure", "source_budget_item_type_code":label,
                "functional_paragraph_code":"|".join(row.get(key, "").strip() for key in (
                    "Codigo_Funcao", "Codigo_Subfuncao", "Codigo_Programa", "Codigo_Acao",
                    "Codigo_Programa_Trabalho",
                )),
                "economic_item_code":composite, "amount_local":value, "currency_code":"BRL",
                "amount_eur":None, "fx_date":None, "is_consolidation_item":False,
                "is_financing":False, "is_summary_row":False, "source_row_number":row_number,
                "source_sheet":"Open_Data_Desp_2025.csv", "source_id":source["id"],
                "ingestion_run_id":source["id"],
                "quality_flags":["official_city_open_data", "maximum_detail_execution", "anchor_only_not_un_agglomeration", "replaces_siconfi_rreo_aggregate_2025"],
                "loaded_at":loaded_at,
            })
    output.mkdir(parents=True, exist_ok=True)
    entities=[{"public_entity_id":"BR:3304557","entity_name":"Município do Rio de Janeiro","entity_type":"municipality","country_code_alpha2":"BR","country_code_alpha3":"BRA","national_entity_code":"3304557","national_entity_code_type":"BR_IBGE_MUNICIPALITY_CODE","is_eu_capital":False,"is_extra_city":True,"default_currency_code":"BRL","eurostat_city_code":None,"eurostat_geography_name":None,"administrative_region_code":"33","administrative_region_name":"Rio de Janeiro","administrative_district_code":None,"administrative_district_name":None,"national_geography_code":"3304557","national_geography_code_type":"BR_IBGE_MUNICIPALITY_CODE","valid_from":None,"valid_to":None,"loaded_at":loaded_at}]
    sources=[{"source_id":source["id"],"public_entity_id":"BR:3304557","source_type":"official_city_open_data","source_name":"Contas Rio Despesas 2025","source_url":source["resource_url"],"dataset_code":"Open_Data_Desp_2025","archive_file":archive.name,"archive_sha256":actual_hash,"retrieved_at":source["retrieved_at"],"notes":source["fiscal_boundary"]+"; landing="+source["landing_url"]+"; dictionary="+source["dictionary_url"]+"; maximum-detail city rows replace overlapping 2025 SICONFI RREO aggregate facts","loaded_at":loaded_at}]
    runs=[{"ingestion_run_id":source["id"],"source_id":source["id"],"started_at":loaded_at,"completed_at":loaded_at,"status":"completed","source_vintage":"2025","source_sha256":actual_hash,"rows_read":rows_read,"rows_loaded":len(facts),"warning_count":1,"error_message":"Anchor municipality only; overlapping 2025 SICONFI RREO aggregate facts must be removed after verified load."}]
    for name,records in (("public_entities",entities),("public_entity_sources",sources),("municipal_budget_line_facts",facts),("ingestion_runs",runs)): write_jsonl(output/f"{name}.jsonl",records)
    manifest={"source_id":source["id"],"source_sha256":actual_hash,"rows_read":rows_read,"facts":len(facts),"entity_count":1,"replacement":{"public_entity_id":"BR:3304557","fiscal_year":2025,"source_id":"br-siconfi-rreo-2025","expected_rows":286}}
    (output/"manifest.json").write_text(json.dumps(manifest,indent=2)+"\n")
    (output/"international_municipal_manifest.json").write_text(json.dumps({"schema_version":"1.0.0","status":"validated","countries":["BRA"],"sources":[source["id"]],"counts":{"public_entities":1,"public_entity_sources":1,"municipal_budget_line_facts":len(facts),"ingestion_runs":1},"source_sha256":actual_hash},indent=2)+"\n")
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-id", required=True)
    parser.add_argument("--config", type=Path, default=CONFIG)
    parser.add_argument("--archive", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    sources = {row["id"]: row for row in json.loads(args.config.read_text(encoding="utf-8"))["sources"]}
    source = sources[args.source_id]
    loaded_at = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    adapter = rio_de_janeiro if source["id"].startswith("br-rio-de-janeiro-") else (lima_mef if source["id"].startswith("pe-lima-mef-") else (sao_paulo if source["id"].startswith("br-sao-paulo-") else (buenos_aires if source["id"].startswith("ar-buenos-aires-") else (san_jose_cr if source["id"].startswith("cr-san-jose-") else (colombia_four_cities if source["id"].startswith("co-four-major-") else (bogota_cuipo if source["id"].startswith("co-bogota-") else (brasilia_df if source["id"].startswith("br-distrito-federal-") else (panama_city if source["id"].startswith("pa-panama-city-") else (guayaquil if source["id"].startswith("ec-guayaquil-") else (mendoza if source["id"].startswith("ar-mendoza-") else montevideo))))))))))
    print(json.dumps(adapter(source, args.archive, args.output, loaded_at)))


if __name__ == "__main__":
    main()
