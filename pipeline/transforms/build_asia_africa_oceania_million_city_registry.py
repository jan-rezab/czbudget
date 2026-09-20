#!/usr/bin/env python3
"""Build the Asia/Africa/Oceania million-city anchor registry from UN WUP F21."""

from __future__ import annotations

import argparse
import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
CONFIG = ROOT / "pipeline/config/asia_africa_oceania_million_cities_sources.json"
PROFILES = ROOT / "data/municipal-expansion"

# Europe and the Americas are owned by separate regional registries. Everything
# else in F21 is in this audit, including Western Asia and transcontinental TUR.
OUT_OF_SCOPE = set("""
ALB AND AUT BEL BGR BIH BLR CHE CYP CZE DEU DNK ESP EST FIN FRA GBR GRC HRV HUN
IRL ISL ITA LIE LTU LUX LVA MCO MDA MKD MLT MNE NLD NOR POL PRT ROU RUS SMR SRB
SVK SVN SWE UKR VAT KOS ARG BHS BLZ BOL BRA BRB CAN CHL COL CRI CUB DOM ECU GTM
GUY HND HTI JAM MEX NIC PAN PER PRI PRY SLV SUR TTO URY USA VEN
""".split())


def digest(path: Path) -> str:
    value = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            value.update(chunk)
    return value.hexdigest()


def build(workbook_path: Path, config_path: Path = CONFIG) -> dict:
    config = json.loads(config_path.read_text(encoding="utf-8"))
    actual_hash = digest(workbook_path)
    if actual_hash != config["un_source"]["sha256"]:
        raise ValueError(f"UN workbook hash mismatch: {actual_hash}")

    sheet = load_workbook(workbook_path, read_only=True, data_only=True)["Data"]
    rows = sheet.iter_rows(values_only=True)
    header = list(next(rows))
    year_index = header.index("2025")
    records = []

    for row in rows:
        population = row[year_index]
        iso3 = row[4]
        if iso3 in OUT_OF_SCOPE or population is None or population < 1000:
            continue
        city_code = str(row[6])
        anchor = config["anchor_entities"].get(city_code)
        record = {
            "un_city_code": int(row[6]),
            "un_city_name": row[7],
            "country": row[1],
            "iso3": iso3,
            "population_2025": int(round(float(population) * 1000)),
            "un_centroid": {"longitude": row[9], "latitude": row[10]},
            "coverage_scope": "UN Degree-of-Urbanization built-up area",
            "coverage_status": "no_published_itemized_anchor",
            "fiscal_scope": None,
            "blocker": config["city_blockers"].get(city_code, "No reviewed published itemized anchor mapping"),
        }
        if anchor:
            profile_path = PROFILES / anchor["country"].lower() / f"{anchor['entity_code']}.json"
            if not profile_path.exists():
                raise ValueError(f"{city_code}: missing profile {profile_path}")
            profile = json.loads(profile_path.read_text(encoding="utf-8"))
            if profile["name"] != anchor["entity_name"]:
                raise ValueError(f"{city_code}: anchor name drift: {profile['name']!r}")
            detail = profile.get("detail", [])
            if not detail:
                raise ValueError(f"{city_code}: profile has no line items")
            record.update({
                "coverage_status": "anchor_only",
                "fiscal_scope": anchor["fiscal_boundary"],
                "profile_entity_code": anchor["entity_code"],
                "profile_entity_name": profile["name"],
                "profile_url": profile["url"],
                "profile_sha256": digest(profile_path),
                "source_url": "https://www.e-stat.go.jp/stat-search/files?toukei=00200251&tstat=000001077755",
                "fiscal_year": 2024,
                "stage": ["actual"],
                "published_detail_rows": len(detail),
                "blocker": "Full agglomeration coverage is not measured; this profile covers only the named city government",
            })
        records.append(record)

    records.sort(key=lambda item: (-item["population_2025"], item["iso3"], item["un_city_name"]))
    status = Counter(item["coverage_status"] for item in records)
    if len(records) != 378:
        raise ValueError(f"Expected 378 in-scope million-plus cities, found {len(records)}")
    return {
        "schema_version": "1.0.0",
        "dataset_id": config["dataset_id"],
        "generated_at": datetime.now(timezone.utc).replace(microsecond=0).isoformat(),
        "source": {**config["un_source"], "retrieved_sha256": actual_hash},
        "methodology": config["methodology"],
        "counts": {
            "in_scope_un_cities": len(records),
            "anchor_only": status["anchor_only"],
            "no_published_itemized_anchor": status["no_published_itemized_anchor"],
            "full_agglomeration": 0,
        },
        "cities": records,
        "priority_official_sources": config["priority_official_sources"],
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--un-workbook", required=True, type=Path)
    parser.add_argument("--config", default=CONFIG, type=Path)
    parser.add_argument("--output", default=ROOT / "data/major-cities/asia-africa-oceania-million-plus.v1.json", type=Path)
    args = parser.parse_args()
    payload = build(args.un_workbook, args.config)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload["counts"], indent=2))


if __name__ == "__main__":
    main()
