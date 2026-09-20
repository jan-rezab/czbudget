#!/usr/bin/env python3
import hashlib, io, json
import openpyxl, requests

URL = "https://einfoszab.budapest.hu/File/DownloadDvD/125496?fileId=201239&filename=szerkesztheto_m01m01_1.+mell%C3%A9klet+-+rendeleti+t%C3%A1bla+2025_rev.xlsx&key=fovarosi-kozgyules-nyilvanos-ulesei&organizationtype=2&parentId=18407&parenttype=2&path=Tev%C3%A9kenys%C3%A9gre%2C+m%C5%B1k%C3%B6d%C3%A9sre+vonatkoz%C3%B3+adatok+%2F+D%C3%B6nt%C3%A9shozatal%2C+%C3%BCl%C3%A9sek+%2F+A+F%C5%91v%C3%A1rosi+K%C3%B6zgy%C5%B1l%C3%A9s+%C3%BCl%C3%A9sei&prevaction=DvdItem&prevcont=Session&sessionType=1&type=5"
r=requests.get(URL,headers={"User-Agent":"Mozilla/5.0 (compatible; PublicSpendingData/1.0)","Referer":"https://einfoszab.budapest.hu/session/DvdItem/125496"},timeout=180)
r.raise_for_status(); body=r.content
wb=openpyxl.load_workbook(io.BytesIO(body),data_only=True,read_only=True)
out={"url":URL,"bytes":len(body),"sha256":hashlib.sha256(body).hexdigest(),"content_type":r.headers.get("content-type"),"sheets":[]}
for ws in wb.worksheets:
    sample=[]
    for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,40),values_only=True):
        vals=[None if v is None else str(v)[:250] for v in row[:18]]
        if any(v not in (None,"") for v in vals): sample.append(vals)
    out["sheets"].append({"title":ws.title,"max_row":ws.max_row,"max_column":ws.max_column,"sample":sample})
print(json.dumps(out,ensure_ascii=False))

