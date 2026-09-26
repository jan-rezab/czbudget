"""Cloud-only immutable acquisition. No warehouse publication in this phase."""
import argparse
import csv
import hashlib
import json
import os
from pathlib import Path
import time
import urllib.request
from google.cloud import storage
import openpyxl
import pyreadstat

PROJECT = 'czbudget-janrezab'
BUCKET = 'czbudget-janrezab-data-layers'

def upload(bucket, name, data):
    blob = bucket.blob(name)
    blob.upload_from_string(data, if_generation_match=0, checksum='auto')
    blob.reload()
    return dict(uri=f'gs://{BUCKET}/{name}', generation=str(blob.generation),
                sha256=hashlib.sha256(data).hexdigest(), bytes=len(data))

def main():
    p=argparse.ArgumentParser();p.add_argument('--loader-sha',required=True);args=p.parse_args()
    release=os.environ['BUILD_ID']; prefix=f'processing-runs/undp/{release}'
    bucket=storage.Client(project=PROJECT).bucket(BUCKET)
    started=time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime())
    sources=[]; previews={}
    for entry in json.loads(Path('pipeline/undp_cloud/sources.json').read_text()):
        data=None
        for attempt in range(4):
            try:
                req=urllib.request.Request(entry['url'],headers={'User-Agent':'PublicSpendingData/1.0 (UNDP warehouse ingestion)'})
                with urllib.request.urlopen(req,timeout=120) as response:
                    data=response.read(100_000_001)
                    if len(data)>100_000_000:raise ValueError('Source exceeded 100 MB bound')
                    content_type=response.headers.get('Content-Type');final_url=response.url
                break
            except Exception:
                if attempt==3:raise
                time.sleep(2**attempt)
        path=Path('/tmp')/f"{entry['id']}.{entry['kind']}";path.write_bytes(data)
        if entry['kind']=='pdf' and not data.startswith(b'%PDF'):raise ValueError('Not a PDF')
        source=dict(entry,content_type=content_type,final_url=final_url,**upload(bucket,f'{prefix}/raw/{path.name}',data));sources.append(source)
        if entry['kind']=='csv':
            with path.open(encoding='utf-8-sig',newline='') as f:
                rows=list(csv.DictReader(f));previews[entry['id']]={'rows':len(rows),'columns':list(rows[0]),'sample':rows[:2]}
        elif entry['kind']=='xlsx':
            wb=openpyxl.load_workbook(path,read_only=True,data_only=True)
            preview={}
            for sheet in wb:
                populated=[]
                for number,row in enumerate(sheet.iter_rows(values_only=True),1):
                    if any(x is not None for x in row):
                        if len(populated)<18:populated.append([number,[str(x) if x is not None else None for x in row]])
                preview[sheet.title]={'rows':sheet.max_row,'columns':sheet.max_column,'first_rows':populated}
            previews[entry['id']]=preview;wb.close()
        elif entry['kind']=='dta':
            df,meta=pyreadstat.read_dta(str(path),metadataonly=True)
            previews[entry['id']]={'rows':meta.number_rows,'columns':meta.column_names,'labels':meta.column_names_to_labels,'variable_value_labels':meta.variable_value_labels}
        print(json.dumps({'event':'raw_source_verified','id':entry['id'],'bytes':len(data),'sha256':source['sha256']}),flush=True)
    manifest=dict(raw_release_id=release,loader_git_sha=args.loader_sha,build_id=release,
                  region='europe-west4',service_account='psd-data-builder@czbudget-janrezab.iam.gserviceaccount.com',
                  started_at=started,completed_at=time.strftime('%Y-%m-%dT%H:%M:%SZ',time.gmtime()),
                  processing_status='raw_verified',publication_status='not_published',sources=sources)
    upload(bucket,f'{prefix}/source-preview.json',(json.dumps(previews,indent=2,default=str)+'\n').encode())
    result=upload(bucket,f'{prefix}/raw-manifest.json',(json.dumps(manifest,indent=2)+'\n').encode())
    print(json.dumps({'event':'raw_acquisition_complete','manifest':result}),flush=True)
if __name__=='__main__':main()
