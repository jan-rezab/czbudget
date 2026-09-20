#!/usr/bin/env python3
"""Acquire, normalize, validate and publish reviewed US city budget sources."""

from __future__ import annotations

import argparse
import csv
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import re
import subprocess
import time
import urllib.parse
import urllib.error
import urllib.request
import zipfile


BUCKET = "gs://czbudget-janrezab-data-layers"
RUN_ROOT = "processing-runs/us-major-cities"
USER_AGENT = "PublicSpendingData/1.0 (official-budget-ingestion)"
SAFE_ID = re.compile(r"^[a-z0-9][a-z0-9._-]{0,119}$")
RETRYABLE_HTTP_STATUS = {429, 500, 502, 503, 504}
REVENUE_ITEM_PREFIXES = frozenset("ABCDTU")
EXPENDITURE_ITEM_PREFIXES = frozenset("EFIJLMS")


def load_registry(path: Path) -> dict:
    registry = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(registry, dict):
        raise ValueError("registry root must be an object")
    return registry


def import_ready_sources(registry: dict) -> list[dict]:
    """Return only explicitly import-ready sources, inheriting city identity."""
    found: list[dict] = []
    for source in registry.get("sources", []):
        if isinstance(source, dict) and source.get("status") == "import_ready":
            found.append(dict(source))
    for city in registry.get("cities", []):
        if not isinstance(city, dict):
            continue
        inherited = {
            "city_slug": city.get("slug") or city.get("city_slug"),
            "city_name": city.get("name") or city.get("city_name"),
            "state": city.get("state"),
            "population": city.get("population_2025") or city.get("population"),
        }
        for source in city.get("sources", []):
            if isinstance(source, dict) and source.get("status") == "import_ready":
                found.append({**inherited, **source})
    for source in found:
        validate_source(source)
    # The registry intentionally references the same national Census archives
    # from every city. Download those once, retaining the applicability list.
    coalesced: dict[str, dict] = {}
    for source in found:
        existing = coalesced.get(source["id"])
        if existing is None:
            source["target_city_slugs"] = ([source["city_slug"]] if source.get("city_slug") else [])
            coalesced[source["id"]] = source
            continue
        shared_keys = ("url", "download_url", "format", "source_kind")
        if source.get("source_kind") != "broad" or any(
            existing.get(key) != source.get(key) for key in shared_keys
        ):
            raise ValueError(f"conflicting import-ready source id: {source['id']}")
        if source.get("city_slug") and source["city_slug"] not in existing["target_city_slugs"]:
            existing["target_city_slugs"].append(source["city_slug"])
        existing["city_slug"] = None
        existing["city_name"] = None
        existing["state"] = None
        existing["population"] = None
    for city in registry.get("cities", []):
        if not isinstance(city, dict):
            continue
        slug = city.get("city_slug") or city.get("slug")
        for source_id in city.get("broad_source_ids", []):
            if source_id in coalesced and slug and slug not in coalesced[source_id]["target_city_slugs"]:
                coalesced[source_id]["target_city_slugs"].append(slug)
    return sorted(coalesced.values(), key=lambda source: source["id"])


def validate_source(source: dict) -> None:
    source_id = source.get("id")
    if not isinstance(source_id, str) or not SAFE_ID.fullmatch(source_id):
        raise ValueError(f"invalid source id: {source_id!r}")
    url = source.get("url") or source.get("download_url")
    if not isinstance(url, str) or not url.startswith("https://"):
        raise ValueError(f"{source_id}: official source URL must use https")
    if source.get("source_kind") not in {"broad", "granular"}:
        raise ValueError(f"{source_id}: source_kind must be broad or granular")
    if source.get("format", "csv") not in {"csv", "json", "jsonl", "arcgis_json", "xlsx", "zip_csv", "zip_fixed_width"}:
        raise ValueError(f"{source_id}: unsupported format")


def gs_uri(run_id: str, suffix: str) -> str:
    if not SAFE_ID.fullmatch(run_id):
        raise ValueError("invalid run id")
    return f"{BUCKET}/{RUN_ROOT}/{run_id}/{suffix}"


def run(command: list[str]) -> None:
    subprocess.run(command, check=True, timeout=1800)


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def urlopen_with_retry(request, *, timeout: int = 180, attempts: int = 5, sleep=time.sleep):
    """Open an official endpoint, retrying only transient transport failures."""
    for attempt in range(1, attempts + 1):
        try:
            return urllib.request.urlopen(request, timeout=timeout)
        except urllib.error.HTTPError as error:
            if error.code not in RETRYABLE_HTTP_STATUS or attempt == attempts:
                raise
        except (urllib.error.URLError, TimeoutError):
            if attempt == attempts:
                raise
        sleep(min(2 ** (attempt - 1), 16))
    raise AssertionError("retry loop exhausted")


def is_socrata(source: dict) -> bool:
    parsed = urllib.parse.urlsplit(source.get("url") or source.get("download_url") or "")
    return source.get("format") == "json" and "/resource/" in parsed.path


def socrata_page_url(url: str, limit: int, offset: int) -> str:
    parsed = urllib.parse.urlsplit(url)
    query = dict(urllib.parse.parse_qsl(parsed.query, keep_blank_values=True))
    query.update({"$limit": str(limit), "$offset": str(offset)})
    query.setdefault("$order", ":id")
    return urllib.parse.urlunsplit((parsed.scheme, parsed.netloc, parsed.path,
                                    urllib.parse.urlencode(query), parsed.fragment))


def arcgis_layer_url(url: str) -> str:
    parsed = urllib.parse.urlsplit(url)
    path = parsed.path.rstrip("/")
    if path.lower().endswith("/query"):
        path = path[:-6]
    return urllib.parse.urlunsplit((parsed.scheme, parsed.netloc, path, "", ""))


def arcgis_page_url(url: str, limit: int, offset: int, order_field: str | None) -> str:
    parsed = urllib.parse.urlsplit(arcgis_layer_url(url) + "/query")
    query = dict(urllib.parse.parse_qsl(urllib.parse.urlsplit(url).query, keep_blank_values=True))
    query.update({"where": "1=1", "outFields": "*", "f": "json",
                  "resultOffset": str(offset), "resultRecordCount": str(limit),
                  "returnGeometry": "false"})
    if order_field:
        query["orderByFields"] = order_field
    return urllib.parse.urlunsplit((parsed.scheme, parsed.netloc, parsed.path,
                                    urllib.parse.urlencode(query), ""))


def arcgis_response(value: object, source_id: str) -> tuple[list[dict], bool]:
    if not isinstance(value, dict):
        raise ValueError(f"{source_id}: ArcGIS response is not an object")
    if value.get("error"):
        error = value["error"]
        raise ValueError(f"{source_id}: ArcGIS service error {error!r}")
    features = value.get("features")
    if not isinstance(features, list):
        raise ValueError(f"{source_id}: ArcGIS response has no features array")
    attributes = []
    for feature in features:
        if not isinstance(feature, dict) or not isinstance(feature.get("attributes"), dict):
            raise ValueError(f"{source_id}: ArcGIS feature lacks attributes")
        attributes.append(feature["attributes"])
    return attributes, bool(value.get("exceededTransferLimit"))


def fetch_socrata(source: dict, destination: Path) -> tuple[str, str, int]:
    """Write a complete, stable Socrata snapshot as one JSON array."""
    url = source.get("url") or source["download_url"]
    limit = int(source.get("page_size", 50_000))
    offset = 0
    total = 0
    content_type = "application/json"
    with destination.open("w", encoding="utf-8") as output:
        output.write("[")
        first = True
        while True:
            page_url = socrata_page_url(url, limit, offset)
            request = urllib.request.Request(page_url, headers={"User-Agent": USER_AGENT})
            with urlopen_with_retry(request) as response:
                content_type = response.headers.get("Content-Type", content_type)
                page = json.load(response)
            if not isinstance(page, list):
                raise ValueError(f"{source['id']}: Socrata page is not an array")
            for record in page:
                if not first:
                    output.write(",")
                json.dump(record, output, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
                first = False
            total += len(page)
            if len(page) < limit:
                break
            offset += limit
        output.write("]\n")
    if total == 0:
        raise ValueError(f"{source['id']}: Socrata returned zero records")
    return url, content_type, total


def fetch_arcgis(source: dict, destination: Path) -> tuple[str, str, int]:
    """Write all ArcGIS feature attributes as a stable JSON array."""
    url = source.get("url") or source["download_url"]
    layer_url = arcgis_layer_url(url)
    metadata_url = layer_url + "?" + urllib.parse.urlencode({"f": "json"})
    request = urllib.request.Request(metadata_url, headers={"User-Agent": USER_AGENT})
    with urlopen_with_retry(request) as response:
        metadata = json.load(response)
    if not isinstance(metadata, dict) or metadata.get("error"):
        raise ValueError(f"{source['id']}: ArcGIS layer metadata error {metadata!r}")
    order_field = source.get("order_by_field") or metadata.get("objectIdField") or metadata.get("objectIdFieldName")
    if not order_field:
        for field in metadata.get("fields", []):
            if isinstance(field, dict) and field.get("type") == "esriFieldTypeOID":
                order_field = field.get("name")
                break
    limit = int(source.get("page_size") or metadata.get("maxRecordCount") or 2_000)
    offset = 0
    total = 0
    with destination.open("w", encoding="utf-8") as output:
        output.write("[")
        first = True
        while True:
            page_url = arcgis_page_url(url, limit, offset, order_field)
            request = urllib.request.Request(page_url, headers={"User-Agent": USER_AGENT})
            with urlopen_with_retry(request) as response:
                page = json.load(response)
            records, exceeded = arcgis_response(page, source["id"])
            for record in records:
                if not first:
                    output.write(",")
                json.dump(record, output, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
                first = False
            total += len(records)
            if not exceeded and len(records) < limit:
                break
            if not records:
                raise ValueError(f"{source['id']}: ArcGIS pagination made no progress")
            offset += len(records)
        output.write("]\n")
    if total == 0:
        raise ValueError(f"{source['id']}: ArcGIS returned zero records")
    return url, "application/json", total


def fetch_one(source: dict, destination: Path) -> dict:
    url = source.get("url") or source["download_url"]
    records = None
    if source.get("format") == "arcgis_json":
        final_url, content_type, records = fetch_arcgis(source, destination)
        size = destination.stat().st_size
    elif is_socrata(source):
        final_url, content_type, records = fetch_socrata(source, destination)
        size = destination.stat().st_size
    else:
        request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
        size = 0
        with urlopen_with_retry(request) as response, destination.open("wb") as output:
            final_url = response.geturl()
            content_type = response.headers.get("Content-Type", "")
            while True:
                block = response.read(1024 * 1024)
                if not block:
                    break
                output.write(block)
                size += len(block)
    if size == 0:
        raise ValueError(f"{source['id']}: empty response")
    return {
        "source_id": source["id"],
        "requested_url": url,
        "resolved_url": final_url,
        "retrieved_at": datetime.now(timezone.utc).isoformat(),
        "sha256": sha256_file(destination),
        "bytes": size,
        "records": records,
        "content_type": content_type,
        "registry_source": source,
    }


def fetch(registry_path: Path, work: Path, run_id: str) -> None:
    sources = import_ready_sources(load_registry(registry_path))
    if not sources:
        raise ValueError("registry has no import_ready sources")
    raw = work / "raw"
    raw.mkdir(parents=True, exist_ok=True)
    receipts = []
    for source in sources:
        extension = {"csv": "csv", "json": "json", "jsonl": "jsonl", "arcgis_json": "json", "xlsx": "xlsx", "zip_csv": "zip",
                     "zip_fixed_width": "zip"}[
            source.get("format", "csv")
        ]
        payload = raw / f"{source['id']}.{extension}"
        receipt = fetch_one(source, payload)
        receipt_path = raw / f"{source['id']}.metadata.json"
        receipt_path.write_text(json.dumps(receipt, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        destination = gs_uri(run_id, f"raw/{payload.name}")
        run(["gsutil", "-h", f"x-goog-meta-sha256:{receipt['sha256']}", "-h",
             f"x-goog-meta-source-id:{source['id']}", "cp", str(payload), destination])
        run(["gsutil", "cp", str(receipt_path), gs_uri(run_id, f"raw/{receipt_path.name}")])
        receipts.append(receipt)
    (work / "fetch-manifest.json").write_text(
        json.dumps({"run_id": run_id, "sources": receipts}, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def nested_records(value, path: str | None):
    if path:
        for part in path.split("."):
            value = value[int(part)] if isinstance(value, list) else value[part]
    if isinstance(value, dict):
        for key in ("data", "results", "records"):
            if isinstance(value.get(key), list):
                return value[key]
    if not isinstance(value, list):
        raise ValueError("JSON source does not resolve to a records array")
    return value


def source_rows(source: dict, payload: Path):
    source_format = source.get("format", "csv")
    if source_format == "csv":
        with payload.open("r", encoding=source.get("encoding", "utf-8-sig"), newline="") as stream:
            yield from csv.DictReader(stream)
    elif source_format in {"json", "arcgis_json"}:
        if (is_socrata(source) or source_format == "arcgis_json") and not source.get("records_path"):
            import ijson
            with payload.open("rb") as stream:
                yield from ijson.items(stream, "item")
        else:
            value = json.loads(payload.read_text(encoding="utf-8"))
            yield from nested_records(value, source.get("records_path"))
    elif source_format == "jsonl":
        with payload.open("r", encoding="utf-8") as stream:
            for line in stream:
                if line.strip():
                    yield json.loads(line)
    elif source_format == "xlsx":
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        workbook = load_workbook(payload, read_only=True, data_only=False)
        reviewed_sheet = source.get("sheet") or source.get("worksheet")
        if reviewed_sheet:
            if reviewed_sheet not in workbook.sheetnames:
                raise ValueError(f"{source['id']}: reviewed sheet is absent")
            sheets = [workbook[reviewed_sheet]]
        else:
            nonempty = []
            for sheet in workbook.worksheets:
                if any(any(value is not None for value in row)
                       for row in sheet.iter_rows(values_only=True)):
                    nonempty.append(sheet)
            if len(nonempty) != 1:
                raise ValueError(f"{source['id']}: expected one non-empty sheet, found {len(nonempty)}")
            sheets = nonempty
        for sheet in sheets:
            for row_number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                if not any(value is not None for value in values):
                    continue
                record = {"_sheet": sheet.title, "_row_number": row_number}
                record.update({get_column_letter(index): value for index, value in enumerate(values, start=1)})
                yield record
        workbook.close()
    elif source_format in {"zip_csv", "zip_fixed_width"}:
        with zipfile.ZipFile(payload) as archive:
            suffixes = (".csv",) if source_format == "zip_csv" else (".txt", ".dat", ".fin")
            csv_names = [name for name in archive.namelist() if name.lower().endswith(suffixes)]
            pattern = source.get("archive_member")
            if pattern:
                csv_names = [name for name in csv_names if re.search(pattern, name)]
            if not csv_names:
                raise ValueError(f"{source['id']}: archive has no matching data members")
            if source_format == "zip_csv" and len(csv_names) != 1:
                raise ValueError(f"{source['id']}: expected one CSV member, found {len(csv_names)}")
            fields = source.get("fixed_width_fields", [])
            for member in sorted(csv_names):
                with archive.open(member) as raw_stream:
                    text_stream = io.TextIOWrapper(raw_stream, encoding=source.get("encoding", "latin-1"))
                    if source_format == "zip_csv":
                        yield from csv.DictReader(text_stream)
                    else:
                        for line_number, line in enumerate(text_stream, start=1):
                            line = line.rstrip("\r\n")
                            if not line:
                                continue
                            record = {"_archive_member": member, "_line_number": line_number,
                                      "_raw_line": line}
                            for field in fields:
                                # Registry positions are zero-based, end-exclusive.
                                record[field["name"]] = line[int(field["start"]):int(field["end"])].strip()
                            yield record


def value_at(row: dict, field_map: dict, name: str):
    source_field = field_map.get(name)
    return row.get(source_field) if source_field else None


def text_at(row: dict, field_map: dict, name: str):
    value = value_at(row, field_map, name)
    return str(value) if value is not None else None


def census_target_cities(registry: dict, source_id: str) -> dict[str, dict]:
    targets = {}
    for city in registry.get("cities", []):
        if source_id not in city.get("broad_source_ids", []):
            continue
        crosswalk = city.get("census_government_id_crosswalk", {})
        government_id = str(crosswalk.get("government_id", ""))
        if not re.fullmatch(r"\d{12}", government_id):
            raise ValueError(f"{city.get('city_slug')}: missing reviewed 12-digit Census government ID")
        if government_id in targets:
            raise ValueError(f"duplicate Census government ID: {government_id}")
        targets[government_id] = city
    if not targets:
        raise ValueError(f"{source_id}: no target cities with reviewed Census government IDs")
    return targets


def census_budget_side(item_code: str) -> str | None:
    if not item_code:
        return None
    if item_code[0] in REVENUE_ITEM_PREFIXES:
        return "revenue"
    if item_code[0] in EXPENDITURE_ITEM_PREFIXES:
        return "expenditure"
    return None


def census_member(archive: zipfile.ZipFile, pattern: str, source_id: str) -> str:
    matches = [name for name in archive.namelist() if re.search(pattern, name)]
    if len(matches) != 1:
        raise ValueError(f"{source_id}: expected one archive member matching {pattern!r}, found {len(matches)}")
    return matches[0]


def parse_census_finance_line(line: str, expected_length: int) -> dict:
    # Census documents a fixed record width, but the published text may omit
    # one final blank flag character. All substantive positions remain fixed.
    if len(line) not in {expected_length - 1, expected_length}:
        raise ValueError(
            f"record has length {len(line)}, expected {expected_length} or {expected_length - 1}"
        )
    item_code = line[12:15]
    return {
        "government_id": line[0:12],
        "item_code": item_code,
        "budget_side": census_budget_side(item_code),
        "amount_local": int(line[15:27].strip() or 0) * 1000,
        "fiscal_year": int(line[27:31]),
        "imputation_flag": line[31:].strip(),
    }


def process_census_finance(source: dict, payload: Path, receipt: dict, registry: dict, output: Path) -> dict:
    """Decode the reviewed Census individual-unit layout and retain only target cities."""
    import pyarrow as pa
    import pyarrow.parquet as pq

    layout = source["census_finance_layout"]
    targets = census_target_cities(registry, source["id"])
    pid_by_id = {}
    with zipfile.ZipFile(payload) as archive:
        pid_name = census_member(archive, layout["pid_member_pattern"], source["id"])
        data_name = census_member(archive, layout["data_member_pattern"], source["id"])
        with archive.open(pid_name) as raw:
            for raw_line in raw:
                line = raw_line.decode(layout.get("encoding", "latin-1")).rstrip("\r\n")
                government_id = line[0:12]
                if government_id in targets:
                    if len(line) < 146:
                        raise ValueError(f"{source['id']}: short PID record for {government_id}")
                    pid_by_id[government_id] = {
                        "census_entity_name": line[12:76].strip(),
                        "pid_population": int(line[116:125].strip() or 0),
                        "pid_population_year": line[125:127].strip(),
                        "fiscal_year_ending": line[140:144].strip(),
                    }
        missing_pid = sorted(set(targets) - set(pid_by_id))
        if missing_pid:
            raise ValueError(f"{source['id']}: target Census IDs absent from PID: {missing_pid}")

        schema = pa.schema([
            ("source_id", pa.string()), ("public_entity_id", pa.string()),
            ("census_government_id", pa.string()), ("city_slug", pa.string()),
            ("city_name", pa.string()), ("state", pa.string()),
            ("census_entity_name", pa.string()), ("pid_population", pa.int64()),
            ("pid_population_year", pa.string()), ("fiscal_year_ending", pa.string()),
            ("fiscal_year", pa.int64()), ("fiscal_period", pa.string()),
            ("reporting_scope", pa.string()), ("budget_stage", pa.string()),
            ("budget_side", pa.string()), ("economic_item_code", pa.string()),
            ("amount_local", pa.int64()), ("currency_code", pa.string()),
            ("imputation_flag", pa.string()), ("is_imputed", pa.bool_()),
            ("source_row_number", pa.int64()), ("source_sheet", pa.string()),
            ("source_line_sha256", pa.string()), ("raw_sha256", pa.string()),
            ("source_url", pa.string()), ("retrieved_at", pa.string()),
            ("coverage_type", pa.string()),
        ])
        writer = pq.ParquetWriter(output, schema, compression="zstd")
        batch = []
        rows = 0
        seen_ids = set()
        skipped_nonbudget = 0
        with archive.open(data_name) as raw:
            for line_number, raw_line in enumerate(raw, start=1):
                line = raw_line.decode(layout.get("encoding", "latin-1")).rstrip("\r\n")
                government_id = line[0:12]
                if government_id not in targets:
                    continue
                expected_length = int(layout["record_length"])
                if len(line) not in {expected_length - 1, expected_length}:
                    raise ValueError(
                        f"{source['id']}: record {line_number} has length {len(line)}, "
                        f"expected {expected_length} or {expected_length - 1}"
                    )
                decoded = parse_census_finance_line(line, expected_length)
                item_code = decoded["item_code"]
                side = decoded["budget_side"]
                if side is None:
                    skipped_nonbudget += 1
                    continue
                fiscal_year = decoded["fiscal_year"]
                if fiscal_year != int(source["fiscal_year"]):
                    raise ValueError(f"{source['id']}: unexpected data year {fiscal_year}")
                imputation_flag = decoded["imputation_flag"]
                city = targets[government_id]
                pid = pid_by_id[government_id]
                batch.append({
                    "source_id": source["id"], "public_entity_id": f"US:{government_id}",
                    "census_government_id": government_id, "city_slug": city["city_slug"],
                    "city_name": city["name"], "state": city["state"], **pid,
                    "fiscal_year": fiscal_year, "fiscal_period": "FY",
                    "reporting_scope": "standalone_accounting_unit", "budget_stage": "actual",
                    "budget_side": side, "economic_item_code": item_code,
                    "amount_local": decoded["amount_local"], "currency_code": "USD",
                    "imputation_flag": imputation_flag, "is_imputed": "I" in imputation_flag,
                    "source_row_number": line_number, "source_sheet": data_name,
                    "source_line_sha256": hashlib.sha256(raw_line.rstrip(b"\r\n")).hexdigest(),
                    "raw_sha256": receipt["sha256"], "source_url": receipt["requested_url"],
                    "retrieved_at": receipt["retrieved_at"],
                    "coverage_type": source.get("coverage_type", "survey"),
                })
                rows += 1
                seen_ids.add(government_id)
                if len(batch) >= 10_000:
                    writer.write_table(pa.Table.from_pylist(batch, schema=schema))
                    batch.clear()
        if batch:
            writer.write_table(pa.Table.from_pylist(batch, schema=schema))
        writer.close()
    missing_data = sorted(set(targets) - seen_ids)
    if missing_data:
        output.unlink(missing_ok=True)
        raise ValueError(f"{source['id']}: target Census IDs have no budget rows: {missing_data}")
    return {
        "source_id": source["id"], "status": "normalized", "rows": rows,
        "entities": len(seen_ids), "expected_entities": len(targets),
        "skipped_nonbudget_rows": skipped_nonbudget, "sha256": sha256_file(output),
        "path": f"normalized/{output.name}",
    }


def process(registry_path: Path, work: Path, run_id: str) -> None:
    import pyarrow as pa
    import pyarrow.parquet as pq

    registry = load_registry(registry_path)
    sources = import_ready_sources(registry)
    normalized = work / "normalized"
    normalized.mkdir(parents=True, exist_ok=True)
    inventories = work / "inventories"
    inventories.mkdir(parents=True, exist_ok=True)
    manifest = json.loads((work / "fetch-manifest.json").read_text(encoding="utf-8"))
    receipt_by_id = {item["source_id"]: item for item in manifest["sources"]}
    output_receipts = []
    for source in sources:
        receipt = receipt_by_id[source["id"]]
        payloads = [path for path in (work / "raw").glob(source["id"] + ".*") if not path.name.endswith(".metadata.json")]
        if len(payloads) != 1:
            raise ValueError(f"{source['id']}: expected one raw payload")
        if sha256_file(payloads[0]) != receipt["sha256"]:
            raise ValueError(f"{source['id']}: raw checksum changed")
        if source.get("census_finance_layout"):
            output = normalized / f"{source['id']}.parquet"
            output_receipts.append(
                process_census_finance(source, payloads[0], receipt, registry, output)
            )
            continue
        if source.get("format") == "zip_fixed_width" and not source.get("fixed_width_fields"):
            with zipfile.ZipFile(payloads[0]) as archive:
                members = [{"name": info.filename, "bytes": info.file_size,
                            "compressed_bytes": info.compress_size, "crc32": f"{info.CRC:08x}"}
                           for info in archive.infolist() if not info.is_dir()]
            if not members:
                raise ValueError(f"{source['id']}: acquired archive has no files")
            output = inventories / f"{source['id']}.json"
            output.write_text(json.dumps({
                "source_id": source["id"], "status": "acquired_only",
                "reason": "No reviewed fixed-width field layout is registered",
                "raw_sha256": receipt["sha256"], "members": members,
            }, indent=2, sort_keys=True) + "\n", encoding="utf-8")
            output_receipts.append({"source_id": source["id"], "status": "acquired_only",
                                    "members": len(members), "sha256": sha256_file(output),
                                    "path": f"inventories/{output.name}"})
            continue
        field_map = source.get("field_map", {})
        output = normalized / f"{source['id']}.parquet"
        schema = pa.schema([
            ("source_id", pa.string()), ("source_kind", pa.string()),
            ("city_slug", pa.string()), ("city_name", pa.string()), ("state", pa.string()),
            ("fiscal_year", pa.string()), ("department", pa.string()), ("fund", pa.string()),
            ("category", pa.string()), ("description", pa.string()), ("amount", pa.string()),
            ("record_ordinal", pa.int64()), ("record_json", pa.string()),
            ("raw_sha256", pa.string()), ("source_url", pa.string()), ("retrieved_at", pa.string()),
        ])
        batch = []
        row_count = 0
        writer = pq.ParquetWriter(output, schema, compression="zstd")
        for ordinal, row in enumerate(source_rows(source, payloads[0]), start=1):
            if not isinstance(row, dict):
                raise ValueError(f"{source['id']}: record {ordinal} is not an object")
            batch.append({
                "source_id": source["id"],
                "source_kind": source["source_kind"],
                "city_slug": source.get("city_slug"),
                "city_name": source.get("city_name"),
                "state": source.get("state"),
                "fiscal_year": str(value_at(row, field_map, "fiscal_year") or source.get("fiscal_year") or ""),
                "department": text_at(row, field_map, "department"),
                "fund": text_at(row, field_map, "fund"),
                "category": text_at(row, field_map, "category"),
                "description": text_at(row, field_map, "description"),
                "amount": text_at(row, field_map, "amount"),
                "record_ordinal": ordinal,
                "record_json": json.dumps(row, ensure_ascii=False, sort_keys=True, default=str),
                "raw_sha256": receipt["sha256"],
                "source_url": receipt["requested_url"],
                "retrieved_at": receipt["retrieved_at"],
            })
            row_count = ordinal
            if len(batch) >= 10_000:
                writer.write_table(pa.Table.from_pylist(batch, schema=schema))
                batch.clear()
        if batch:
            writer.write_table(pa.Table.from_pylist(batch, schema=schema))
        writer.close()
        if not row_count:
            output.unlink(missing_ok=True)
            raise ValueError(f"{source['id']}: parsed zero records")
        output_receipts.append({"source_id": source["id"], "status": "normalized",
                                "rows": row_count, "sha256": sha256_file(output),
                                "path": f"normalized/{output.name}"})
    (work / "validation.json").write_text(
        json.dumps({"run_id": run_id, "status": "validated", "outputs": output_receipts},
                   indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )


def publish(registry_path: Path, work: Path, run_id: str) -> None:
    validation = json.loads((work / "validation.json").read_text(encoding="utf-8"))
    if validation.get("status") != "validated" or not validation.get("outputs"):
        raise ValueError("validated outputs are required before publication")
    for item in validation["outputs"]:
        path = work / item["path"]
        if sha256_file(path) != item["sha256"]:
            raise ValueError(f"{item['source_id']}: derived checksum changed")
        run(["gsutil", "-h", f"x-goog-meta-sha256:{item['sha256']}", "cp", str(path),
             gs_uri(run_id, item["path"])])
    registry_hash = sha256_file(registry_path)
    run(["gsutil", "cp", str(work / "fetch-manifest.json"), gs_uri(run_id, "fetch-manifest.json")])
    run(["gsutil", "cp", str(work / "validation.json"), gs_uri(run_id, "validation.json")])
    completed = work / "completed.json"
    completed.write_text(json.dumps({
        "run_id": run_id,
        "status": "completed",
        "completed_at": datetime.now(timezone.utc).isoformat(),
        "registry_sha256": registry_hash,
        "source_status_counts": {
            "normalized": sum(item["status"] == "normalized" for item in validation["outputs"]),
            "acquired_only": sum(item["status"] == "acquired_only" for item in validation["outputs"]),
        },
        "outputs": validation["outputs"],
    }, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    # This is deliberately the final write: consumers must ignore runs without it.
    run(["gsutil", "-h", "Cache-Control:no-store", "cp", str(completed), gs_uri(run_id, "completed.json")])


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("command", choices=("fetch", "process", "publish"))
    parser.add_argument("--registry", type=Path, required=True)
    parser.add_argument("--work", type=Path, required=True)
    parser.add_argument("--run-id", required=True)
    args = parser.parse_args()
    {"fetch": fetch, "process": process, "publish": publish}[args.command](
        args.registry, args.work, args.run_id
    )


if __name__ == "__main__":
    main()
