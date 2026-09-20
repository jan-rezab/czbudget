#!/usr/bin/env python3
"""Cloud ingestion for official Canadian million-city municipal budget lines."""
import csv, hashlib, io, json, re
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
import requests
from google.cloud import bigquery
from openpyxl import load_workbook

PROJECT="czbudget-janrezab"; DATASET="budget_detail"
SOURCES={
 "ca-toronto-approved-operating-2025": {"entity":"CA:3520005","url":"https://ckan0.cf.opendata.inter.prod-toronto.ca/dataset/2c90a5d3-5598-4c02-abf2-169456c8f1f1/resource/f9def3c1-a97f-4d31-bc58-c0494d750b80/download/approved-operating-budget-summary-2025.xlsx"},
 "ca-montreal-approved-operating-2025": {"entity":"CA:2466023","url":"https://donnees.montreal.ca/dataset/2b7dcae7-a3e5-4f5d-81de-ccde2c518e55/resource/1db52976-a013-49ac-a133-27b55585556a/download/budget-fonctionnement-2025.xlsx"},
 "ca-vancouver-approved-capital-2025": {"entity":"CA:5915022","url":"https://opendata.vancouver.ca/api/explore/v2.1/catalog/datasets/2025-multi-year-capital-project-budget-requests-and-capital-expenditure-budget/exports/csv?lang=en&timezone=America%2FVancouver&use_labels=true&delimiter=%2C"},
}

def number(v):
 try: return Decimal(str(v))
 except (InvalidOperation,TypeError): return None

def fact(entity, source, rowno, side, function, item, value, loaded, scope="standalone_municipality"):
 value=abs(value).quantize(Decimal("0.000000001"))
 return {"public_entity_id":entity,"fiscal_year":2025,"fiscal_period":"FY","reporting_scope":scope,"budget_stage":"enacted","budget_side":side,
 "source_budget_item_type_code":item,"functional_paragraph_code":function,"economic_item_code":item,"functional_classification_id":f"{entity}_FUNCTION_2025","economic_classification_id":f"{entity}_ECONOMIC_2025",
 "amount_local":str(value),"currency_code":"CAD","amount_eur":None,"fx_date":None,"is_consolidation_item":False,"is_financing":False,"is_summary_row":False,
 "source_row_number":rowno,"source_sheet":source,"source_id":source,"ingestion_run_id":source+"-v1","coverage_type":"published_subset","is_imputed":False,
 "quality_flags":["official_municipal_open_data","anchor_legal_government_only"],"loaded_at":loaded}

def toronto(body,loaded):
 wb=load_workbook(io.BytesIO(body),read_only=True,data_only=True); ws=wb["Open Data"]; it=ws.iter_rows(values_only=True); headers=list(next(it)); out=[]
 for n,row in enumerate(it,2):
  rec=dict(zip(headers,row)); val=number(rec.get("2025")); side="revenue" if rec.get("Expense/Revenue")=="Revenues" else "expenditure"
  if val is None or not val: continue
  function=" > ".join(str(rec.get(k) or "Unspecified") for k in ("Program","Service","Activity")); item=" > ".join(str(rec.get(k) or "Unspecified") for k in ("Category Name","Sub-Category Name","Commitment item"))
  out.append(fact(SOURCES["ca-toronto-approved-operating-2025"]["entity"],"ca-toronto-approved-operating-2025",n,side,function,item,val,loaded))
 return out

def montreal(body,loaded):
 wb=load_workbook(io.BytesIO(body),read_only=True,data_only=True); ws=wb["SOMM CM"]; out=[]; side=None
 for n,row in enumerate(ws.iter_rows(min_row=4,values_only=True),4):
  label=str(row[0] or "").strip(); upper=label.upper()
  if upper=="REVENUS": side="revenue"; continue
  if "DÉPENSE" in upper: side="expenditure"; continue
  val=number(row[7] if len(row)>7 else None)
  if side and val is not None and val and not upper.startswith("TOTAL"):
   # The workbook publishes values in thousands of CAD. Round after scaling
   # to remove binary-float noise introduced by XLSX numeric cells.
   cad=(val*1000).quantize(Decimal("0.01"))
   out.append(fact(SOURCES["ca-montreal-approved-operating-2025"]["entity"],"ca-montreal-approved-operating-2025",n,side,"municipal_budget",label,cad,loaded))
 return out

def vancouver(body,loaded):
 reader=csv.DictReader(io.StringIO(body.decode("utf-8-sig"))); out=[]; col="Annual Capital Expenditure-2025 Capital Expenditure Budget"
 for n,row in enumerate(reader,2):
  val=number(row.get(col));
  if val is None or not val: continue
  function=" > ".join(str(row.get(k) or "Unspecified") for k in ("Service Category 1","Service Category 2","Service Category 3")); item=str(row.get("Project/Program Name") or "Unspecified")
  # The portal occasionally serializes calculated currency values with a
  # binary-float tail (for example, 706167.8200000001).
  out.append(fact(SOURCES["ca-vancouver-approved-capital-2025"]["entity"],"ca-vancouver-approved-capital-2025",n,"expenditure",function,item,val.quantize(Decimal("0.01")),loaded))
 return out

def main():
 loaded=datetime.now(timezone.utc).isoformat(); facts=[]; receipts={}
 parsers={"ca-toronto-approved-operating-2025":toronto,"ca-montreal-approved-operating-2025":montreal,"ca-vancouver-approved-capital-2025":vancouver}
 for sid,src in SOURCES.items():
  r=requests.get(src["url"],timeout=120); r.raise_for_status(); body=r.content; rows=parsers[sid](body,loaded)
  if not rows: raise RuntimeError(f"zero rows: {sid}")
  facts.extend(rows); receipts[sid]={"url":src["url"],"sha256":hashlib.sha256(body).hexdigest(),"bytes":len(body),"entity_id":src["entity"],"rows_loaded":len(rows),"year":2025,"stage":"enacted","currency":"CAD"}
 client=bigquery.Client(project=PROJECT); target=f"{PROJECT}.{DATASET}.municipal_budget_line_facts"; stage=f"{PROJECT}.{DATASET}._canada_city_budget_stage"
 client.query(f"CREATE OR REPLACE TABLE `{stage}` LIKE `{target}`").result(); client.load_table_from_json(facts,stage).result()
 if client.get_table(stage).num_rows!=len(facts): raise RuntimeError("stage count mismatch")
 ids=','.join("'"+x["entity"]+"'" for x in SOURCES.values())
 client.query(f"BEGIN TRANSACTION; DELETE FROM `{target}` WHERE fiscal_year=2025 AND public_entity_id IN ({ids}) AND STARTS_WITH(source_id,'ca-'); INSERT INTO `{target}` SELECT * FROM `{stage}` WHERE fiscal_year=2025; COMMIT TRANSACTION; DROP TABLE `{stage}`;").result()
 receipt={"retrieved_at":loaded,"fiscal_scope":"Anchor legal municipalities only; not UN built-up areas","sources":receipts,"rows_loaded":len(facts)}
 Path('/workspace/canada-receipt.json').write_text(json.dumps(receipt,indent=2)+'\n'); print(json.dumps(receipt))
if __name__=='__main__': main()
