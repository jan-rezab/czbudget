import hashlib, io, json
from pathlib import Path
import requests
from openpyxl import load_workbook

SOURCES = {
 "toronto-operating-2025": "https://ckan0.cf.opendata.inter.prod-toronto.ca/dataset/2c90a5d3-5598-4c02-abf2-169456c8f1f1/resource/f9def3c1-a97f-4d31-bc58-c0494d750b80/download/approved-operating-budget-summary-2025.xlsx",
 "montreal-operating-2025": "https://donnees.montreal.ca/dataset/2b7dcae7-a3e5-4f5d-81de-ccde2c518e55/resource/1db52976-a013-49ac-a133-27b55585556a/download/budget-fonctionnement-2025.xlsx",
 "vancouver-capital-2025": "https://opendata.vancouver.ca/api/explore/v2.1/catalog/datasets/2025-multi-year-capital-project-budget-requests-and-capital-expenditure-budget/exports/csv?lang=en&timezone=America%2FVancouver&use_labels=true&delimiter=%2C",
}

def main():
 out={}
 for sid,url in SOURCES.items():
  r=requests.get(url,timeout=120); r.raise_for_status(); body=r.content
  item={"url":url,"bytes":len(body),"sha256":hashlib.sha256(body).hexdigest()}
  if url.endswith("xlsx"):
   wb=load_workbook(io.BytesIO(body),read_only=True,data_only=True)
   item["sheets"]={}
   for ws in wb.worksheets:
    rows=[]
    for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,30),values_only=True):
     values=[None if v is None else str(v)[:160] for v in row]
     if any(v not in (None,"") for v in values): rows.append(values[:20])
    item["sheets"][ws.title]={"max_row":ws.max_row,"max_column":ws.max_column,"sample":rows}
  else:
   item["sample"]=body[:12000].decode("utf-8-sig",errors="replace").splitlines()[:20]
  out[sid]=item
 try:
  r=requests.get("https://api.us.socrata.com/api/catalog/v1",params={"q":"budget","search_context":"data.calgary.ca"},timeout=60)
  out["calgary-catalog"]={"status":r.status_code,"sample":r.text[:20000]}
 except Exception as e: out["calgary-catalog"]={"error":str(e)}
 Path('/workspace/canada-discovery.json').write_text(json.dumps(out,ensure_ascii=False,indent=2)+'\n')
 print(json.dumps({k:{x:y for x,y in v.items() if x in ('bytes','sha256','status','error')} for k,v in out.items()}))
if __name__=='__main__': main()
