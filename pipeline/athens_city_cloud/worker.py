#!/usr/bin/env python3
"""Load coded 2024 Athens municipal budget-execution rows from official XLSX files."""
import hashlib
import io
import json
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import requests
from google.cloud import bigquery

PROJECT = "czbudget-janrezab"
DATASET = "budget_detail"
ENTITY = "GR:ATHENS"
MUNICIPALITY = "ΔΗΜΟΣ ΑΘΗΝΑΙΩΝ"
RUN = "gr-athens-municipal-execution-2024-v1"
SOURCES = {
    "gr-moi-municipal-execution-revenue-2024": {
        "side": "revenue",
        "url": "https://data.gov.gr/dataset/5554b44b-9e1a-4c84-81fc-0e956b391642/resource/87d6067c-57f4-42f9-b4e9-da84801ec50c/download/01.dimoi_stoixeia-ektelesis-proypologismoy_esoda_2024.xlsx",
        "stages": {9: "enacted", 10: "revised", 8: "actual"},
    },
    "gr-moi-municipal-execution-expenditure-2024": {
        "side": "expenditure",
        "url": "https://data.gov.gr/dataset/5554b44b-9e1a-4c84-81fc-0e956b391642/resource/7e53135e-2a01-43c1-9780-53e3067379fc/download/02.dimoi_stoixeia-ektelesis-proypologismoy_exoda_2024.xlsx",
        "stages": {8: "enacted", 9: "revised", 11: "actual"},
    },
}


def money(value):
    try:
        amount = Decimal(str(value))
    except (InvalidOperation, TypeError):
        return None
    return amount if amount.is_finite() else None


def normalize(body, source_id, source, loaded_at):
    book = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    ws = book[book.sheetnames[0]]
    rows = []
    source_lines = 0
    for row_number, values in enumerate(ws.iter_rows(values_only=True), 1):
        if row_number == 1 or str(values[3]).strip().upper() != MUNICIPALITY:
            continue
        if int(values[4]) != 2024 or int(values[5]) != 12:
            continue
        code = str(values[6]).strip()
        label = str(values[7]).strip()
        if not code or code.lower() == "none" or not label or label.lower() == "none":
            continue
        source_lines += 1
        for column, stage in source["stages"].items():
            amount = money(values[column])
            if amount is None or amount == 0:
                continue
            formatted = format(abs(amount), "f")
            rows.append({
                "public_entity_id": ENTITY, "fiscal_year": 2024, "fiscal_period": "FY",
                "reporting_scope": "standalone_municipality", "budget_stage": stage,
                "budget_side": source["side"], "source_budget_item_type_code": label,
                "functional_paragraph_code": None, "economic_item_code": f"{source['side']}:{code}",
                "functional_classification_id": None,
                "economic_classification_id": "GR_KAE_FOURTH_GRADE_2024",
                "amount_local": formatted, "currency_code": "EUR", "amount_eur": formatted,
                "fx_date": "2024-12-31", "is_consolidation_item": False,
                "is_financing": False, "is_summary_row": False, "source_row_number": row_number,
                "source_sheet": ws.title, "source_id": source_id, "ingestion_run_id": RUN,
                "coverage_type": "official_coded_municipal_execution_fourth_grade",
                "is_imputed": False,
                "quality_flags": ["official_greek_interior_ministry_open_data", "legal_municipality_anchor_only", "month_12_full_year", "published_fourth_grade_code", f"source_column:{stage}"],
                "loaded_at": loaded_at,
            })
    if source_lines < 20 or len(rows) < 50:
        raise RuntimeError(f"Athens source contract failed for {source_id}: lines={source_lines}, facts={len(rows)}")
    return rows, source_lines, ws.title


def main():
    loaded_at = datetime.now(timezone.utc).isoformat()
    session = requests.Session()
    session.headers["User-Agent"] = "Mozilla/5.0 (compatible; PublicSpendingData/1.0)"
    facts = []
    receipt_sources = {}
    for source_id, source in SOURCES.items():
        response = session.get(source["url"], timeout=180)
        response.raise_for_status()
        parsed, source_lines, sheet = normalize(response.content, source_id, source, loaded_at)
        facts.extend(parsed)
        receipt_sources[source_id] = {
            "official_url": source["url"], "sha256": hashlib.sha256(response.content).hexdigest(),
            "bytes": len(response.content), "source_lines": source_lines,
            "rows_loaded": len(parsed), "source_sheet": sheet,
        }
    client = bigquery.Client(project=PROJECT)
    target = f"{PROJECT}.{DATASET}.municipal_budget_line_facts"
    stage = f"{PROJECT}.{DATASET}._athens_city_budget_stage"
    client.query(f"CREATE OR REPLACE TABLE `{stage}` LIKE `{target}`").result()
    client.load_table_from_json(facts, stage).result()
    if client.get_table(stage).num_rows != len(facts):
        raise RuntimeError("Athens staging row count mismatch")
    client.query(f"""
      BEGIN TRANSACTION;
      DELETE FROM `{target}` WHERE public_entity_id = '{ENTITY}' AND fiscal_year = 2024
        AND STARTS_WITH(source_id, 'gr-moi-municipal-execution-');
      INSERT INTO `{target}` SELECT * FROM `{stage}` WHERE fiscal_year = 2024;
      COMMIT TRANSACTION;
      DROP TABLE `{stage}`;
    """).result()
    counts = [dict(row) for row in client.query(f"""
      SELECT budget_side, budget_stage, COUNT(*) AS row_count,
             CAST(SUM(amount_local) AS STRING) AS amount_local
      FROM `{target}`
      WHERE public_entity_id = '{ENTITY}' AND fiscal_year = 2024
        AND STARTS_WITH(source_id, 'gr-moi-municipal-execution-')
      GROUP BY budget_side, budget_stage ORDER BY budget_side, budget_stage
    """).result()]
    verified = sum(int(row["row_count"]) for row in counts)
    if verified != len(facts):
        raise RuntimeError(f"Athens warehouse count mismatch: expected={len(facts)}, verified={verified}")
    receipt = {
        "status": "warehouse_loaded", "retrieved_at": loaded_at, "entity_id": ENTITY,
        "fiscal_year": 2024,
        "scope": "Municipality of Athens legal government only; excludes Attica Region, neighboring municipalities, and the UN built-up city.",
        "publisher": "Hellenic Ministry of Interior, Local Government Performance Monitoring Hub",
        "catalog_url": "https://data.gov.gr/dataset/oikonomika-stoicheia-ota-etos-2024-komvos-parakoloythisis-epidoseon-topikis-aytodi",
        "rows_loaded": len(facts), "warehouse_verification": counts, "sources": receipt_sources,
    }
    Path("/workspace/athens-receipt.json").write_text(json.dumps(receipt, ensure_ascii=False, indent=2) + "\n")
    print(json.dumps(receipt, ensure_ascii=False))


if __name__ == "__main__":
    main()
