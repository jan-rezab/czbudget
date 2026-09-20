#!/usr/bin/env python3
import hashlib
import io
import json

import openpyxl
import requests

SOURCES = {
    "revenue": "https://data.gov.gr/dataset/5554b44b-9e1a-4c84-81fc-0e956b391642/resource/87d6067c-57f4-42f9-b4e9-da84801ec50c/download/01.dimoi_stoixeia-ektelesis-proypologismoy_esoda_2024.xlsx",
    "expenditure": "https://data.gov.gr/dataset/5554b44b-9e1a-4c84-81fc-0e956b391642/resource/7e53135e-2a01-43c1-9780-53e3067379fc/download/02.dimoi_stoixeia-ektelesis-proypologismoy_exoda_2024.xlsx",
}
s = requests.Session()
s.headers["User-Agent"] = "Mozilla/5.0 (compatible; PublicSpendingData/1.0)"
out = {}
for side, url in SOURCES.items():
    r = s.get(url, timeout=120)
    r.raise_for_status()
    book = openpyxl.load_workbook(io.BytesIO(r.content), read_only=True, data_only=True)
    sheets = []
    for ws in book.worksheets:
        preview = []
        athens_rows = []
        for index, row in enumerate(ws.iter_rows(values_only=True), 1):
            values = [value.isoformat() if hasattr(value, "isoformat") else value for value in row]
            if index <= 12:
                preview.append(values)
            if "ΑΘΗΝ" in " ".join(str(value).upper() for value in values if value is not None):
                athens_rows.append({"row": index, "values": values})
                if len(athens_rows) >= 8:
                    break
        sheets.append({"title": ws.title, "max_row": ws.max_row, "max_column": ws.max_column, "preview": preview, "athens_rows": athens_rows})
    out[side] = {"url": url, "bytes": len(r.content), "sha256": hashlib.sha256(r.content).hexdigest(), "sheets": sheets}
print(json.dumps(out, ensure_ascii=False, default=str))
