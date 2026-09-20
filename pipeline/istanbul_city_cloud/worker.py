#!/usr/bin/env python3
"""Load official İBB 2025 enacted and executed final-account leaf lines."""
import hashlib
import io
import json
import zipfile
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
import requests
from google.cloud import bigquery

PROJECT="czbudget-janrezab"; DATASET="budget_detail"; ENTITY="TR:IBB:46.34.01"
URL="https://uploads.ibb.istanbul/uploads/2025_yili_kesin_hesap_a7055163c7.zip"
EXPECTED="c174f5c7b36f04ae0495bc52b858b3b0c88c3a1847c1908d35f4702fc822c9fb"
RUN="tr-istanbul-final-account-2025-v1"

def dec(value):
    try:
        out=Decimal(str(value)); return out if out.is_finite() and out != 0 else None
    except (InvalidOperation,ValueError,TypeError): return None

def fact(side,stage,code,label,amount,row_number,sheet,loaded,flags):
    return {"public_entity_id":ENTITY,"fiscal_year":2025,"fiscal_period":"FY","reporting_scope":"istanbul_metropolitan_municipality_final_account","budget_stage":stage,"budget_side":side,"source_budget_item_type_code":label,"functional_paragraph_code":code if side=="expenditure" else "UNSPECIFIED","economic_item_code":code,"functional_classification_id":"TR_MAHALLI_FONK_FIN_2025" if side=="expenditure" else "TR_MAHALLI_REVENUE_NOT_FUNCTIONAL","economic_classification_id":"TR_MAHALLI_EKO_2025","amount_local":str(abs(amount)),"currency_code":"TRY","amount_eur":None,"fx_date":None,"is_consolidation_item":False,"is_financing":False,"is_summary_row":False,"source_row_number":row_number,"source_sheet":sheet,"source_id":"tr-ibb-final-account-2025","ingestion_run_id":RUN,"coverage_type":"official_coded_final_account_leaf","is_imputed":False,"quality_flags":["official_ibb_final_account","legal_metropolitan_municipality_anchor_only","leaf_classification_only",*flags],"loaded_at":loaded}

def parse_expense(payload,loaded):
    book=openpyxl.load_workbook(io.BytesIO(payload),read_only=True,data_only=True); sheet=book.worksheets[0]; rows=[]
    hierarchy={0:("", ""),2:("", ""),4:("", ""),6:("", ""),8:("", ""),10:("", "")}
    for n,values in enumerate(sheet.iter_rows(min_row=5,values_only=True),start=5):
        for idx in hierarchy:
            if values[idx] not in (None,""):
                hierarchy[idx]=(str(values[idx]).strip(),str(values[idx+1] or "").strip())
                for deeper in hierarchy:
                    if deeper>idx: hierarchy[deeper]=("","")
        leaf=str(values[12] or "").strip()
        if not leaf: continue
        label=str(values[13] or "").strip(); daire=hierarchy[0][0]; f1=hierarchy[2][0]; f2=hierarchy[4][0]; f3=hierarchy[6][0]; fin=hierarchy[8][0]; e1=hierarchy[10][0]
        code=f"{daire}.{f1}.{f2}.{f3}.{fin}.{e1}.{leaf}"
        full_label=" | ".join(x for x in (hierarchy[0][1],label) if x)
        for stage,index in (("enacted",15),("actual",20)):
            amount=dec(values[index])
            if amount is not None: rows.append(fact("expenditure",stage,code,full_label,amount,n,sheet.title,loaded,[f"daire:{daire}"]))
    return rows

def parse_revenue(payload,loaded):
    book=openpyxl.load_workbook(io.BytesIO(payload),read_only=True,data_only=True); sheet=book.worksheets[0]; rows=[]
    for n,values in enumerate(sheet.iter_rows(min_row=6,values_only=True),start=6):
        codes=[str(values[i] or "").strip() for i in range(5)]
        if not codes[4]: continue
        code=".".join(codes); label=str(values[5] or "").strip()
        for stage,index in (("enacted",7),("actual",15)):
            amount=dec(values[index])
            if amount is not None: rows.append(fact("revenue",stage,code,label,amount,n,sheet.title,loaded,[]))
    return rows

def main():
    loaded=datetime.now(timezone.utc).isoformat(); response=requests.get(URL,timeout=240,headers={"User-Agent":"Mozilla/5.0 (compatible; PublicSpendingData/1.0)"}); response.raise_for_status()
    digest=hashlib.sha256(response.content).hexdigest()
    if digest!=EXPECTED: raise RuntimeError(f"archive hash drift: {digest}")
    archive=zipfile.ZipFile(io.BytesIO(response.content)); members={x.filename:archive.read(x.filename) for x in archive.infolist()}
    expense=next((v for k,v in members.items() if k.startswith("1-") and k.endswith(".xlsx")),None); revenue=next((v for k,v in members.items() if k.startswith("2-") and k.endswith(".xlsx")),None)
    if not expense or not revenue: raise RuntimeError("final-account flow workbooks missing")
    facts=parse_expense(expense,loaded)+parse_revenue(revenue,loaded)
    if len(facts)<300: raise RuntimeError(f"too few leaf-stage facts: {len(facts)}")
    client=bigquery.Client(project=PROJECT); target=f"{PROJECT}.{DATASET}.municipal_budget_line_facts"; stage=f"{PROJECT}.{DATASET}._istanbul_final_account_stage"
    client.query(f"CREATE OR REPLACE TABLE `{stage}` LIKE `{target}`").result(); client.load_table_from_json(facts,stage).result()
    if client.get_table(stage).num_rows!=len(facts): raise RuntimeError("stage row mismatch")
    client.query(f"BEGIN TRANSACTION; DELETE FROM `{target}` WHERE public_entity_id='{ENTITY}' AND fiscal_year=2025 AND source_id='tr-ibb-final-account-2025'; INSERT INTO `{target}` SELECT * FROM `{stage}` WHERE fiscal_year=2025; COMMIT TRANSACTION; DROP TABLE `{stage}`;").result()
    counts={}
    for row in facts: counts[f"{row['budget_stage']}:{row['budget_side']}"]=counts.get(f"{row['budget_stage']}:{row['budget_side']}",0)+1
    receipt={"retrieved_at":loaded,"entity_id":ENTITY,"entity_name":"İstanbul Büyükşehir Belediye Başkanlığı","fiscal_year":2025,"reporting_scope":"İBB legal metropolitan municipality only; districts, affiliated entities and full UN built-up area are not consolidated","official_landing_page":"https://ibb.istanbul/ibb/butce-ve-yatirimlar/","official_url":URL,"archive_sha256":digest,"archive_bytes":len(response.content),"member_sha256":{k:hashlib.sha256(v).hexdigest() for k,v in members.items()},"budget_stages":["enacted","actual"],"rows_loaded":len(facts),"rows_by_stage_side":counts,"excluded_components":["financing classification","balance sheet","trial balance"],"ingestion_run_id":RUN}
    Path("/workspace/istanbul-receipt.json").write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+"\n"); print(json.dumps(receipt,ensure_ascii=False))
if __name__=="__main__": main()
