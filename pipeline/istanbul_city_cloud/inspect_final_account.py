#!/usr/bin/env python3
import hashlib
import io
import json
import zipfile

import openpyxl
import requests

URL = "https://uploads.ibb.istanbul/uploads/2025_yili_kesin_hesap_a7055163c7.zip"
response = requests.get(URL, timeout=240, headers={"User-Agent": "Mozilla/5.0 (compatible; PublicSpendingData/1.0)"})
response.raise_for_status()
archive = zipfile.ZipFile(io.BytesIO(response.content))
result = {"url": URL, "bytes": len(response.content), "sha256": hashlib.sha256(response.content).hexdigest(), "members": []}
for info in archive.infolist():
    member = {"name": info.filename, "bytes": info.file_size, "crc": info.CRC}
    if info.filename.lower().endswith((".xlsx", ".xlsm")):
        payload = archive.read(info.filename)
        book = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        member["sha256"] = hashlib.sha256(payload).hexdigest()
        member["sheets"] = []
        for sheet in book.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = [value for value in row if value is not None]
                if values:
                    rows.append(values[:20])
                if len(rows) >= 12:
                    break
            member["sheets"].append({"title": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column, "sample": rows})
            if info.filename.startswith(("1-", "2-")) and sheet == book.worksheets[0]:
                member["raw_rows"] = [
                    [cell.value for cell in row[:25]]
                    for row in sheet.iter_rows(min_row=1, max_row=15)
                ]
    result["members"].append(member)
print(json.dumps(result, ensure_ascii=False, default=str))
